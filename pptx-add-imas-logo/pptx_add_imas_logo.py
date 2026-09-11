# /// script
# dependencies = [
#   "google-auth",
#   "google-api-python-client",
#   "python-pptx",
# ]
# ///

import socket
# Force IPv4 to bypass broken IPv6 resolution on macOS
orig_getaddrinfo = socket.getaddrinfo
def forced_getaddrinfo(*args, **kwargs):
    args = list(args)
    if len(args) > 2:
        args[2] = socket.AF_INET
    else:
        while len(args) < 3:
            args.append(None)
        args[2] = socket.AF_INET
    return orig_getaddrinfo(*args, **kwargs)
socket.getaddrinfo = forced_getaddrinfo

import os
import sys
import uuid
import argparse
from google.oauth2 import service_account
from googleapiclient.discovery import build

# URL to the uploaded logo image (publicly accessible on GitHub raw)
DEFAULT_IMAGE_URL = 'https://raw.githubusercontent.com/sunlight3d/skills/master/pptx-add-imas-logo/logo_imas.png'
LOCAL_LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'logo_imas.png')

# Dimensions for IMAS logo (1527 x 688) -> width 2,000,000 EMU
IMG_WIDTH_EMU = 2000000
IMG_HEIGHT_EMU = int(IMG_WIDTH_EMU * (688 / 1527))  # 901,113 EMU
MARGIN = 100000

DEFAULT_CREDS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'credentials.json')
if not os.path.exists(DEFAULT_CREDS):
    for candidate in [
        r'C:\code\skills\mcp-google-sheets\credentials.json',
        '/Users/hoangnd/Documents/funix-auto-sheet-f464a0b5957e.json',
        '/Users/hoangnd/Documents/connect-gemini-api-471309-f8bf3d1a89f5.json',
    ]:
        if os.path.exists(candidate):
            DEFAULT_CREDS = candidate
            break

# -------------------------------------------------------------
# Google Slides Online Mode
# -------------------------------------------------------------
def process_google_slides(slides_service, presentation_id, image_url, remove_only=False):
    try:
        presentation = slides_service.presentations().get(presentationId=presentation_id).execute()
    except Exception as e:
        print(f"  Error fetching presentation {presentation_id}: {e}")
        return
        
    slides = presentation.get('slides', [])
    page_size = presentation.get('pageSize', {})
    slide_width = page_size.get('width', {}).get('magnitude', 16256000)
    slide_height = page_size.get('height', {}).get('magnitude', 9144000)

    pos_x = slide_width - IMG_WIDTH_EMU - MARGIN
    pos_y = slide_height - IMG_HEIGHT_EMU - MARGIN

    requests = []
    for slide in slides:
        slide_id = slide.get('objectId')
        
        # Delete old logos and cover rectangles created by this script
        elements = slide.get('pageElements', [])
        for el in elements:
            el_id = el.get('objectId', '')
            if el_id.startswith('cover_img_') or el_id.startswith('cover_rect_'):
                requests.append({'deleteObject': {'objectId': el_id}})
                
        if not remove_only:
            rect_id = f"cover_rect_{slide_id}_{uuid.uuid4().hex[:8]}"
            img_id = f"cover_img_{slide_id}_{uuid.uuid4().hex[:8]}"
            
            # 1. Create a white rectangle to hide the old logo/watermark
            requests.append({
                'createShape': {
                    'objectId': rect_id,
                    'shapeType': 'RECTANGLE',
                    'elementProperties': {
                        'pageObjectId': slide_id,
                        'size': {
                            'width': {'magnitude': IMG_WIDTH_EMU, 'unit': 'EMU'},
                            'height': {'magnitude': IMG_HEIGHT_EMU, 'unit': 'EMU'}
                        },
                        'transform': {
                            'scaleX': 1.0,
                            'scaleY': 1.0,
                            'translateX': pos_x,
                            'translateY': pos_y,
                            'unit': 'EMU'
                        }
                    }
                }
            })
            
            # Format shape with white solid fill, no outline
            requests.append({
                'updateShapeProperties': {
                    'objectId': rect_id,
                    'shapeProperties': {
                        'shapeBackgroundFill': {
                            'solidFill': {
                                'color': {
                                    'rgbColor': {
                                        'red': 1.0,
                                        'green': 1.0,
                                        'blue': 1.0
                                    }
                                }
                            }
                        },
                        'outline': {
                            'propertyState': 'NOT_RENDERED'
                        }
                    },
                    'fields': 'shapeBackgroundFill.solidFill.color,outline'
                }
            })
            
            # 2. Add the IMAS logo on top
            requests.append({
                'createImage': {
                    'objectId': img_id,
                    'url': image_url,
                    'elementProperties': {
                        'pageObjectId': slide_id,
                        'size': {
                            'width': {'magnitude': IMG_WIDTH_EMU, 'unit': 'EMU'},
                            'height': {'magnitude': IMG_HEIGHT_EMU, 'unit': 'EMU'}
                        },
                        'transform': {
                            'scaleX': 1.0,
                            'scaleY': 1.0,
                            'translateX': pos_x,
                            'translateY': pos_y,
                            'unit': 'EMU'
                        }
                    }
                }
            })

    if requests:
        action_name = "Removing" if remove_only else "Adding"
        print(f"  {action_name} IMAS logo on {len(slides)} slides...")
        body = {'requests': requests}
        slides_service.presentations().batchUpdate(presentationId=presentation_id, body=body).execute()
        print("  Done!")
    else:
        print("  No slides found in the presentation.")

def get_presentations_in_folder(drive_service, folder_id):
    query = f"'{folder_id}' in parents and mimeType='application/vnd.google-apps.presentation' and trashed=false"
    results = []
    page_token = None
    while True:
        response = drive_service.files().list(q=query, spaces='drive', fields='nextPageToken, files(id, name)', pageToken=page_token).execute()
        results.extend(response.get('files', []))
        page_token = response.get('nextPageToken', None)
        if page_token is None:
            break
    return results

# -------------------------------------------------------------
# Local PPTX Mode
# -------------------------------------------------------------
def process_local_pptx(file_path, output_path=None, logo_path=None):
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.dml.color import RGBColor

    if not os.path.exists(file_path):
        print(f"Error: File not found at {file_path}")
        sys.exit(1)

    logo_file = logo_path if (logo_path and os.path.exists(logo_path)) else LOCAL_LOGO_PATH
    if not os.path.exists(logo_file):
        print(f"Error: Logo file not found at {logo_file}")
        sys.exit(1)

    prs = Presentation(file_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    rect_width = Emu(IMG_WIDTH_EMU)
    rect_height = Emu(IMG_HEIGHT_EMU)
    margin_emu = Emu(MARGIN)

    pos_x = slide_width - rect_width - margin_emu
    pos_y = slide_height - rect_height - margin_emu

    for slide in prs.slides:
        # 1. Add white cover rectangle
        cover_shape = slide.shapes.add_shape(
            1, # MSO_SHAPE.RECTANGLE
            pos_x, pos_y, rect_width, rect_height
        )
        cover_shape.fill.solid()
        cover_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        cover_shape.line.fill.background()
        cover_shape.line.width = 0

        # 2. Add IMAS logo image on top
        slide.shapes.add_picture(
            logo_file,
            pos_x, pos_y, rect_width, rect_height
        )

    save_path = output_path if output_path else file_path
    prs.save(save_path)
    print(f"Success! Added IMAS logo to {len(prs.slides)} slides. Saved to {save_path}")

# -------------------------------------------------------------
# Main CLI Entry Point
# -------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Add IMAS logo to the bottom right of Google Slides or local PPTX to cover previous logos.")
    
    # Check if first arg is a known command or fallback to single presentation / local file
    parser.add_argument("target", help="The Google Slides presentation ID / Folder ID, or path to local .pptx file")
    parser.add_argument("--folder", action="store_true", help="Treat target as Google Drive Folder ID")
    parser.add_argument("--local", action="store_true", help="Treat target as local .pptx file")
    parser.add_argument("--credentials", default=DEFAULT_CREDS,
                        help="Path to Google Service Account Credentials JSON")
    parser.add_argument("--image-url", default=DEFAULT_IMAGE_URL,
                        help="URL of the logo image to insert for Google Slides")
    parser.add_argument("--logo-path", default=LOCAL_LOGO_PATH,
                        help="Local path to logo_imas.png for local PPTX")
    parser.add_argument("--output", default=None,
                        help="Output path for local .pptx file (defaults to overwrite in-place)")
    parser.add_argument("--remove-only", action="store_true",
                        help="Remove existing logos from Google Slides without adding new ones")
    
    args = parser.parse_args()

    # Determine execution mode: local PPTX vs Google Slides
    is_local = args.local or args.target.lower().endswith('.pptx') or os.path.isfile(args.target)
    
    if is_local:
        print(f"Processing local PowerPoint presentation: {args.target}")
        process_local_pptx(args.target, args.output, args.logo_path)
    else:
        # Google Slides mode
        scopes = ['https://www.googleapis.com/auth/presentations', 'https://www.googleapis.com/auth/drive.readonly']
        creds = service_account.Credentials.from_service_account_file(args.credentials, scopes=scopes)
        slides_service = build('slides', 'v1', credentials=creds)

        if args.folder:
            drive_service = build('drive', 'v3', credentials=creds)
            print(f"Searching for presentations in Google Drive folder {args.target}...")
            presentations = get_presentations_in_folder(drive_service, args.target)
            print(f"Found {len(presentations)} presentations.")
            for i, p in enumerate(presentations):
                print(f"[{i+1}/{len(presentations)}] Processing '{p['name']}' (ID: {p['id']})")
                process_google_slides(slides_service, p['id'], args.image_url, args.remove_only)
        else:
            print(f"Processing single Google Slides presentation (ID: {args.target})")
            process_google_slides(slides_service, args.target, args.image_url, args.remove_only)

if __name__ == '__main__':
    main()
