# /// script
# dependencies = [
#   "google-auth",
#   "google-api-python-client",
#   "python-pptx",
#   "requests",
# ]
# ///

import socket
# Force IPv4 to bypass broken IPv6 resolution on macOS
orig_getaddrinfo = socket.getaddrinfo
def forced_getaddrinfo(*args, **kwargs):
    args = list(args)
    if len(args) > 2:
        args[2] = socket.AF_INET
    else:
        while len(args) < 3:
            args.append(None)
        args[2] = socket.AF_INET
    return orig_getaddrinfo(*args, **kwargs)
socket.getaddrinfo = forced_getaddrinfo

import os
import re
import sys
import uuid
import argparse
from pathlib import Path
from google.oauth2 import service_account
from googleapiclient.discovery import build

# URL to the uploaded logo image (publicly accessible 16:9 high resolution)
DEFAULT_IMAGE_URL = 'https://raw.githubusercontent.com/sunlight3d/skills/master/aptech-fill-logo-benh-vien-quang-xuong/logo_4552x2560.jpg'
FALLBACK_IMAGE_URL = 'https://files.catbox.moe/1emza2.jpg'
LOCAL_LOGO_NAME = 'logo_4552x2560.jpg'

def get_script_dir():
    return Path(__file__).resolve().parent

def parse_input_target(input_str):
    """
    Parses input string to determine if it is:
    - A Google Slides presentation URL or ID
    - A Google Drive folder URL or ID
    - A local .pptx file path
    - A local directory containing .pptx files
    """
    s = input_str.strip()
    
    # 1. Local path check
    if os.path.exists(s):
        if os.path.isdir(s):
            return 'local_folder', s
        elif s.lower().endswith('.pptx'):
            return 'local_file', s

    # 2. Google Drive Folder URL
    # Format: https://drive.google.com/drive/folders/<FOLDER_ID> or drive/u/0/folders/<FOLDER_ID>
    folder_match = re.search(r'drive\.google\.com/drive/(?:u/\d+/)?folders/([a-zA-Z0-9_-]+)', s)
    if folder_match:
        return 'drive_folder', folder_match.group(1)
        
    # 3. Google Slides URL
    # Format: https://docs.google.com/presentation/d/<PRESENTATION_ID>/...
    slide_match = re.search(r'docs\.google\.com/presentation/d/([a-zA-Z0-9_-]+)', s)
    if slide_match:
        return 'google_slide', slide_match.group(1)

    # 4. Fallback: Raw ID
    return 'raw_id', s

def process_google_slides(slides_service, presentation_id, image_url, remove_only=False):
    """Inserts a white cover rectangle and overlays BV Quang Xuong logo at the bottom right."""
    try:
        presentation = slides_service.presentations().get(presentationId=presentation_id).execute()
    except Exception as e:
        print(f"  ❌ Error fetching presentation {presentation_id}: {e}")
        return False
        
    title = presentation.get('title', presentation_id)
    slides = presentation.get('slides', [])
    page_size = presentation.get('pageSize', {})
    slide_width = page_size.get('width', {}).get('magnitude', 16256000)
    slide_height = page_size.get('height', {}).get('magnitude', 9144000)

    # Logo dimensions: 4552 x 2560 (ratio 16:9)
    # Width 2,000,000 EMU covers typical NotebookLM logo completely (~157 pt)
    IMG_WIDTH_EMU = 2000000
    IMG_HEIGHT_EMU = int(IMG_WIDTH_EMU * (2560 / 4552))  # ~1,124,780 EMU

    MARGIN_X = 50000  # 50k EMU padding from right
    MARGIN_Y = 40000  # 40k EMU padding from bottom
    pos_x = slide_width - IMG_WIDTH_EMU - MARGIN_X
    pos_y = slide_height - IMG_HEIGHT_EMU - MARGIN_Y

    # White cover rectangle (slightly larger to ensure 100% watermark occlusion)
    RECT_WIDTH_EMU = IMG_WIDTH_EMU + 20000
    RECT_HEIGHT_EMU = IMG_HEIGHT_EMU + 20000
    rect_x = pos_x - 10000
    rect_y = pos_y - 10000

    requests = []
    
    for slide in slides:
        slide_id = slide.get('objectId')
        elements = slide.get('pageElements', [])
        
        # 1. Clean up previous covers or logos created by this skill
        for el in elements:
            el_id = el.get('objectId', '')
            if (el_id.startswith('cover_qx_') or 
                el_id.startswith('cover_img_') or 
                el_id.startswith('cover_rect_')):
                requests.append({'deleteObject': {'objectId': el_id}})

        if not remove_only:
            rect_id = f"cover_qx_rect_{slide_id}_{uuid.uuid4().hex[:6]}"
            img_id = f"cover_qx_img_{slide_id}_{uuid.uuid4().hex[:6]}"

            # A. Insert pure white rectangle with no border
            requests.append({
                'createShape': {
                    'objectId': rect_id,
                    'shapeType': 'RECTANGLE',
                    'elementProperties': {
                        'pageObjectId': slide_id,
                        'size': {
                            'width': {'magnitude': RECT_WIDTH_EMU, 'unit': 'EMU'},
                            'height': {'magnitude': RECT_HEIGHT_EMU, 'unit': 'EMU'}
                        },
                        'transform': {
                            'scaleX': 1,
                            'scaleY': 1,
                            'translateX': rect_x,
                            'translateY': rect_y,
                            'unit': 'EMU'
                        }
                    }
                }
            })
            requests.append({
                'updateShapeProperties': {
                    'objectId': rect_id,
                    'shapeProperties': {
                        'shapeBackgroundFill': {
                            'solidFill': {
                                'color': {
                                    'rgbColor': {'red': 1.0, 'green': 1.0, 'blue': 1.0}
                                }
                            }
                        },
                        'outline': {
                            'propertyState': 'NOT_RENDERED'
                        }
                    },
                    'fields': 'shapeBackgroundFill,outline'
                }
            })

            # B. Insert BV Quang Xuong Logo image on top
            requests.append({
                'createImage': {
                    'objectId': img_id,
                    'url': image_url,
                    'elementProperties': {
                        'pageObjectId': slide_id,
                        'size': {
                            'width': {'magnitude': IMG_WIDTH_EMU, 'unit': 'EMU'},
                            'height': {'magnitude': IMG_HEIGHT_EMU, 'unit': 'EMU'}
                        },
                        'transform': {
                            'scaleX': 1,
                            'scaleY': 1,
                            'translateX': pos_x,
                            'translateY': pos_y,
                            'unit': 'EMU'
                        }
                    }
                }
            })

    if requests:
        action_name = "Đang gỡ logo" if remove_only else "Đang chèn logo BV ĐK Quảng Xương vào"
        print(f"  {action_name} {len(slides)} slides của presentation '{title}'...")
        body = {'requests': requests}
        slides_service.presentations().batchUpdate(presentationId=presentation_id, body=body).execute()
        print(f"  ✅ Hoàn thành presentation: '{title}' ({len(slides)} slides)")
        return True
    else:
        print(f"  ℹ️ Không có thao tác nào cần thực hiện cho: '{title}'")
        return True

def get_presentations_in_folder(drive_service, folder_id, recursive=True):
    """Finds all Google Slides presentations inside a Google Drive folder (supports recursive)."""
    presentations = []
    folders_to_scan = [folder_id]
    scanned_folders = set()

    while folders_to_scan:
        current_folder = folders_to_scan.pop(0)
        if current_folder in scanned_folders:
            continue
        scanned_folders.add(current_folder)

        # 1. Get presentations in current folder
        query = f"'{current_folder}' in parents and mimeType='application/vnd.google-apps.presentation' and trashed=false"
        page_token = None
        while True:
            response = drive_service.files().list(
                q=query, 
                spaces='drive', 
                fields='nextPageToken, files(id, name)', 
                pageToken=page_token
            ).execute()
            presentations.extend(response.get('files', []))
            page_token = response.get('nextPageToken', None)
            if page_token is None:
                break

        # 2. Get subfolders if recursive
        if recursive:
            folder_query = f"'{current_folder}' in parents and mimeType='application/vnd.google-apps.folder' and trashed=false"
            page_token = None
            while True:
                response = drive_service.files().list(
                    q=folder_query,
                    spaces='drive',
                    fields='nextPageToken, files(id, name)',
                    pageToken=page_token
                ).execute()
                for subf in response.get('files', []):
                    if subf['id'] not in scanned_folders:
                        folders_to_scan.append(subf['id'])
                page_token = response.get('nextPageToken', None)
                if page_token is None:
                    break

    return presentations

def process_local_pptx(pptx_path, logo_img_path, output_path=None):
    """Processes local PowerPoint file using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation(pptx_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    # 2,000,000 EMU width
    img_width = Emu(2000000)
    img_height = Emu(int(2000000 * (2560 / 4552)))
    margin_x = Emu(50000)
    margin_y = Emu(40000)
    
    pos_x = slide_width - img_width - margin_x
    pos_y = slide_height - img_height - margin_y

    for idx, slide in enumerate(prs.slides):
        # 1. Add white background rectangle
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            pos_x - Emu(10000), 
            pos_y - Emu(10000), 
            img_width + Emu(20000), 
            img_height + Emu(20000)
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(255, 255, 255)
        rect.line.fill.background()

        # 2. Add logo image
        slide.shapes.add_picture(
            logo_img_path, 
            pos_x, 
            pos_y, 
            width=img_width, 
            height=img_height
        )

    save_path = output_path if output_path else pptx_path
    prs.save(save_path)
    print(f"  ✅ Đã xử lý và lưu file local PPTX: {save_path} ({len(prs.slides)} slides)")
    return True

def find_default_credentials():
    script_dir = get_script_dir()
    candidates = [
        script_dir / 'credentials.json',
        Path('/Users/hoangnd/.gemini/config/skills/pptx-add-qnet-logo/credentials.json'),
        Path('/Users/hoangnd/Documents/connect-gemini-api-471309-2da0973af1ba.json'),
        Path('/Users/hoangnd/Documents/connect-gemini-api-471309-f8bf3d1a89f5.json'),
        Path('/Users/hoangnd/Documents/funix-auto-sheet-f464a0b5957e.json'),
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    return None

def main():
    parser = argparse.ArgumentParser(
        description="Chèn logo BV ĐK Quảng Xương (16:9) vào góc dưới bên phải Google Slides để che NotebookLM."
    )
    parser.add_argument("target", help="Link Google Slide, ID presentation, Link thư mục Google Drive, hoặc file/thư mục PPTX")
    parser.add_argument("--folder", action="store_true", help="Chỉ định mục tiêu là Thư mục Google Drive chứa nhiều slide")
    parser.add_argument("--credentials", default=None, help="Đường dẫn file JSON credentials của Service Account")
    parser.add_argument("--image-url", default=DEFAULT_IMAGE_URL, help="URL ảnh logo công khai (mặc định dùng ảnh 16:9 chất lượng cao)")
    parser.add_argument("--remove-only", action="store_true", help="Chỉ gỡ các logo đã chèn trước đó, không thêm mới")
    parser.add_argument("--output", default=None, help="Đường dẫn lưu file PPTX đầu ra (áp dụng cho local pptx)")

    args = parser.parse_args()

    # Determine credentials
    creds_path = args.credentials or find_default_credentials()
    
    target_type, target_val = parse_input_target(args.target)
    if args.folder:
        target_type = 'drive_folder'

    # Local file / directory execution
    if target_type == 'local_file':
        local_logo = str(get_script_dir() / LOCAL_LOGO_NAME)
        if not os.path.exists(local_logo):
            local_logo = 'Baitap02/logo_4552x2560.jpg'
        print(f"Xử lý file PowerPoint cục bộ: {target_val}")
        process_local_pptx(target_val, local_logo, args.output)
        return

    elif target_type == 'local_folder':
        local_logo = str(get_script_dir() / LOCAL_LOGO_NAME)
        if not os.path.exists(local_logo):
            local_logo = 'Baitap02/logo_4552x2560.jpg'
        pptx_files = list(Path(target_val).glob('*.pptx'))
        print(f"Tìm thấy {len(pptx_files)} file .pptx trong thư mục: {target_val}")
        for p in pptx_files:
            process_local_pptx(str(p), local_logo)
        return

    # Google Slides / Drive execution
    if not creds_path or not os.path.exists(creds_path):
        print(f"❌ Lỗi: Không tìm thấy file credentials Google Service Account! Vui lòng cung cấp qua --credentials")
        sys.exit(1)

    scopes = [
        'https://www.googleapis.com/auth/presentations',
        'https://www.googleapis.com/auth/drive.readonly'
    ]
    creds = service_account.Credentials.from_service_account_file(creds_path, scopes=scopes)
    slides_service = build('slides', 'v1', credentials=creds)

    if target_type == 'drive_folder':
        drive_service = build('drive', 'v3', credentials=creds)
        print(f"🔍 Đang tìm các file Google Slides trong thư mục Drive ID: {target_val}...")
        presentations = get_presentations_in_folder(drive_service, target_val)
        print(f"👉 Tìm thấy {len(presentations)} bài thuyết trình trong thư mục.")
        
        for i, p in enumerate(presentations):
            print(f"\n[{i+1}/{len(presentations)}] Đang xử lý: '{p['name']}' (ID: {p['id']})")
            process_google_slides(slides_service, p['id'], args.image_url, args.remove_only)
        print("\n🎉 Tất cả bài thuyết trình trong thư mục đã được xử lý xong!")
    else:
        # Single presentation
        print(f"🚀 Đang xử lý Google Slide đơn (ID: {target_val})...")
        success = process_google_slides(slides_service, target_val, args.image_url, args.remove_only)
        if success:
            print("🎉 Hoàn tất chèn logo Bệnh viện ĐK Quảng Xương!")

if __name__ == '__main__':
    main()
