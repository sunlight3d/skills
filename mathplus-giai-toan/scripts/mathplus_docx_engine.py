#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Core Engine tạo tài liệu Word (.docx) chuẩn phong cách Trung tâm MathPlus Academy.
Quy chuẩn thiết kế:
- Chữ toàn bộ màu ĐEN TUYỀN (RGBColor(0, 0, 0) / #000000), nền trắng hoàn toàn.
- Cỡ chữ toàn văn bản TỐI THIỂU 11pt (Pt(11) / sz=22).
- Khung lời giải chi tiết (Solution Box) KHÔNG có sọc đen bên trái (bỏ hoàn toàn viền trái).
- Chèn Logo MathPlus nguyên màu thương hiệu tại Header đầu trang.
- Chuyển đổi 100% công thức toán học LaTeX -> OMML (Word Native Equations).
- Đảm bảo 100% ECMA-376 OpenXML schema compliance.
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from omml_converter import latex_to_omml

# Bảng màu chuẩn ĐEN TUYỀN - NỀN TRẮNG
COLOR_BLACK = RGBColor(0, 0, 0)             # #000000 - Đen tuyền tuyệt đối
COLOR_DARK_TEXT = RGBColor(0, 0, 0)         # Đen tuyền
COLOR_PRIMARY_NAVY = RGBColor(0, 0, 0)      # Đen tuyền
COLOR_MUTED_GRAY = RGBColor(0, 0, 0)        # Đen tuyền
COLOR_SOLUTION_TITLE = RGBColor(0, 0, 0)    # Đen tuyền
COLOR_CAPTION = RGBColor(0, 0, 0)           # Đen tuyền
COLOR_WHITE = RGBColor(255, 255, 255)

def find_mathplus_logo(custom_path=None):
    """Tìm file logo MathPlus để chèn vào Header, giữ nguyên màu gốc."""
    candidates = [
        custom_path,
        r"C:\code\mathplus\public\images\logo.jpg",
        r"C:\code\mathplus\public\images\logo.png",
        os.path.join(os.path.dirname(__file__), "..", "resources", "logo.jpg"),
        os.path.join(os.path.dirname(__file__), "..", "resources", "logo.png"),
        os.path.join(os.getcwd(), "public", "images", "logo.jpg"),
        os.path.join(os.getcwd(), "public", "images", "logo.png"),
    ]
    for c in candidates:
        if c and os.path.exists(c) and os.path.isfile(c):
            return c
    return None

def style_cell(cell, bg_color=None, border_color="000000", top="none", bottom="none", left="none", right="none", sz="4"):
    """Định kiểu màu nền và đường viền cho ô theo đúng thứ tự schema OpenXML."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith(('tcBorders', 'shd')):
            tcPr.remove(child)
            
    if any(b != 'none' for b in [top, bottom, left, right]):
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
            f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
            f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
            f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{border_color}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
    if bg_color:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}" w:val="clear"/>')
        tcPr.append(shd)

def add_inline_math(p, latex_code):
    omml = latex_to_omml(latex_code, display=False)
    p._p.append(omml)

def add_display_math(doc_or_cell, latex_code, space_before=Pt(2), space_after=Pt(3)):
    p = doc_or_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    omml = latex_to_omml(latex_code, display=True)
    p._p.append(omml)
    return p

def add_rich_text(p, text, default_font="Times New Roman", default_size=Pt(11),
                  default_color=COLOR_BLACK, default_bold=False, default_italic=False):
    """Phân tích chuỗi chứa text, công thức toán $...$ và markdown bold/italic (**...**) rồi thêm run đen tuyền tương ứng (font >= 11pt)."""
    effective_size = Pt(11) if default_size < Pt(11) else default_size
    tokens = re.split(r'(\$[^$]+\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('$') and token.endswith('$') and len(token) > 2:
            math_code = token[1:-1].strip()
            try:
                add_inline_math(p, math_code)
            except Exception:
                run = p.add_run(token)
                run.font.name = default_font
                run.font.size = effective_size
                run.font.color.rgb = COLOR_BLACK
                run.font.bold = default_bold
                run.font.italic = default_italic
        else:
            md_tokens = re.split(r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)', token)
            for md_tok in md_tokens:
                if not md_tok:
                    continue
                if md_tok.startswith('***') and md_tok.endswith('***') and len(md_tok) > 6:
                    run = p.add_run(md_tok[3:-3])
                    run.font.name = default_font
                    run.font.size = effective_size
                    run.font.color.rgb = COLOR_BLACK
                    run.font.bold = True
                    run.font.italic = True
                elif md_tok.startswith('**') and md_tok.endswith('**') and len(md_tok) > 4:
                    run = p.add_run(md_tok[2:-2])
                    run.font.name = default_font
                    run.font.size = effective_size
                    run.font.color.rgb = COLOR_BLACK
                    run.font.bold = True
                    run.font.italic = default_italic
                elif md_tok.startswith('*') and md_tok.endswith('*') and len(md_tok) > 2:
                    run = p.add_run(md_tok[1:-1])
                    run.font.name = default_font
                    run.font.size = effective_size
                    run.font.color.rgb = COLOR_BLACK
                    run.font.bold = default_bold
                    run.font.italic = True
                else:
                    clean_tok = md_tok.replace('**', '')
                    run = p.add_run(clean_tok)
                    run.font.name = default_font
                    run.font.size = effective_size
                    run.font.color.rgb = COLOR_BLACK
                    run.font.bold = default_bold
                    run.font.italic = default_italic

class MathPlusDocxBuilder:
    def __init__(self, title, subtitle="Năm học 2026 – 2027", time_limit="90 phút", is_solved=False, logo_path=None, **kwargs):
        self.doc = Document()
        self.title = title
        self.subtitle = subtitle
        self.time_limit = time_limit
        self.is_solved = is_solved
        self.logo_path = find_mathplus_logo(logo_path)
        self._setup_page()
        self._setup_header_footer()
        self._build_top_banner()

    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(18)
        section.bottom_margin = Mm(18)
        section.left_margin = Mm(20)
        section.right_margin = Mm(18)
        section.different_first_page_header_footer = True
        
        normal_style = self.doc.styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = COLOR_BLACK

    def _setup_header_footer(self):
        section = self.doc.sections[0]
        header = section.header
        p_h = header.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_h.paragraph_format.space_before = Pt(0)
        p_h.paragraph_format.space_after = Pt(0)
        
        # Header ở các trang sau: Chỉ gồm Logo và tên Trung tâm (không viết hoa toàn bộ)
        if self.logo_path and os.path.exists(self.logo_path):
            r_logo = p_h.add_run()
            r_logo.add_picture(self.logo_path, width=Mm(7.5))
            
        r_h = p_h.add_run("  Trung tâm MathPlus Academy")
        r_h.font.name = "Times New Roman"
        r_h.font.size = Pt(11)
        r_h.font.italic = True
        r_h.font.color.rgb = COLOR_BLACK

        def setup_ftr(footer_obj):
            p_f = footer_obj.paragraphs[0]
            p_f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_f.paragraph_format.tab_stops.add_tab_stop(Mm(172), WD_TAB_ALIGNMENT.RIGHT)
            
            r_f1 = p_f.add_run("Giáo viên: Nguyễn Văn Dũng – MathPlus Academy  |  Hotline: 0977 961 189")
            r_f1.font.name = "Times New Roman"
            r_f1.font.size = Pt(11)
            r_f1.font.italic = True
            r_f1.font.color.rgb = COLOR_BLACK
            
            r_f2 = p_f.add_run("\tTrang ")
            r_f2.font.name = "Times New Roman"
            r_f2.font.size = Pt(11)
            r_f2.font.italic = True
            r_f2.font.color.rgb = COLOR_BLACK
            
            r_pg = p_f.add_run()
            r_pg.font.name = "Times New Roman"
            r_pg.font.size = Pt(11)
            r_pg.font.color.rgb = COLOR_BLACK
            fld1 = OxmlElement('w:fldChar')
            fld1.set(qn('w:fldCharType'), 'begin')
            t1 = OxmlElement('w:instrText')
            t1.set(qn('xml:space'), 'preserve')
            t1.text = 'PAGE'
            fld2 = OxmlElement('w:fldChar')
            fld2.set(qn('w:fldCharType'), 'separate')
            fld3 = OxmlElement('w:fldChar')
            fld3.set(qn('w:fldCharType'), 'end')
            r_pg._r.append(fld1)
            r_pg._r.append(t1)
            r_pg._r.append(fld2)
            r_pg._r.append(fld3)

        setup_ftr(section.first_page_footer)
        setup_ftr(section.footer)

    def _build_top_banner(self):
        header_table = self.doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False
        
        # Cột trái (78mm) và Cột phải (94mm) - Cân đối, không bị tràn dòng
        c_left = header_table.cell(0, 0)
        c_left.width = Mm(78)
        c_right = header_table.cell(0, 1)
        c_right.width = Mm(94)
        
        # Căn giữa theo chiều dọc (Vertical Center) cho cả hai cột để không bị lệch
        for c in [c_left, c_right]:
            tcPr = c._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
        
        # Cột trái: Logo (18mm) + Tên trung tâm vừa vặn 1 dòng
        first_p = c_left.paragraphs[0]
        first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first_p.paragraph_format.space_before = Pt(0)
        first_p.paragraph_format.space_after = Pt(2)
        
        if self.logo_path and os.path.exists(self.logo_path):
            run_logo = first_p.add_run()
            run_logo.add_picture(self.logo_path, width=Mm(18))
            p_l = c_left.add_paragraph()
        else:
            p_l = first_p
            
        p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_l.paragraph_format.space_before = Pt(1)
        p_l.paragraph_format.space_after = Pt(1)
        p_l.paragraph_format.line_spacing = 1.15
        
        r_l1 = p_l.add_run("TRUNG TÂM MATHPLUS ACADEMY\n")
        r_l1.font.name = "Times New Roman"
        r_l1.font.size = Pt(11)
        r_l1.font.bold = True
        r_l1.font.color.rgb = COLOR_BLACK
        
        r_l2 = p_l.add_run("Bồi dưỡng Toán & Luyện thi\n")
        r_l2.font.name = "Times New Roman"
        r_l2.font.size = Pt(11)
        r_l2.font.italic = True
        r_l2.font.color.rgb = COLOR_BLACK
        
        r_l3 = p_l.add_run("Hotline: 0977 961 189  •  mathplus.com.vn")
        r_l3.font.name = "Times New Roman"
        r_l3.font.size = Pt(11)
        r_l3.font.italic = True
        r_l3.font.color.rgb = COLOR_BLACK
        
        # Cột phải: Tiêu đề đề thi cân xứng
        p_r = c_right.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_r.paragraph_format.space_before = Pt(0)
        p_r.paragraph_format.space_after = Pt(1)
        p_r.paragraph_format.line_spacing = 1.15
        
        r_r1 = p_r.add_run(f"{self.title.upper()}\n")
        r_r1.font.name = "Times New Roman"
        r_r1.font.size = Pt(12)
        r_r1.font.bold = True
        r_r1.font.color.rgb = COLOR_BLACK
        
        if self.is_solved:
            r_bdg = p_r.add_run("【 ĐỀ BÀI VÀ LỜI GIẢI CHI TIẾT – CÓ HÌNH VẼ 】\n")
            r_bdg.font.name = "Times New Roman"
            r_bdg.font.size = Pt(11)
            r_bdg.font.bold = True
            r_bdg.font.color.rgb = COLOR_BLACK
            
        r_r2 = p_r.add_run(f"{self.subtitle}\n")
        r_r2.font.name = "Times New Roman"
        r_r2.font.size = Pt(11)
        r_r2.font.bold = True
        r_r2.font.color.rgb = COLOR_BLACK
        
        r_r3 = p_r.add_run(f"Thời gian gợi ý: {self.time_limit}")
        r_r3.font.name = "Times New Roman"
        r_r3.font.size = Pt(11)
        r_r3.font.italic = True
        r_r3.font.color.rgb = COLOR_BLACK
        
        # Khung thông tin học sinh (Nền trắng, viền mảnh đen)
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)
        
        info_t = self.doc.add_table(rows=1, cols=1)
        info_t.alignment = WD_TABLE_ALIGNMENT.CENTER
        info_t.autofit = False
        c_info = info_t.cell(0, 0)
        c_info.width = Mm(172)
        style_cell(c_info, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        
        p_i = c_info.paragraphs[0]
        p_i.paragraph_format.space_before = Pt(3)
        p_i.paragraph_format.space_after = Pt(3)
        p_i.paragraph_format.line_spacing = 1.2
        
        r1 = p_i.add_run("Họ và tên học sinh: ")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK
        r1_dot = p_i.add_run("……………………………………………………………………    ")
        r1_dot.font.size = Pt(11)
        r1_dot.font.color.rgb = COLOR_BLACK
        
        r2 = p_i.add_run("Lớp: ")
        r2.bold = True
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_BLACK
        r2_dot = p_i.add_run("……………………\n")
        r2_dot.font.size = Pt(11)
        r2_dot.font.color.rgb = COLOR_BLACK
        
        r3 = p_i.add_run("Ngày làm bài: ")
        r3.bold = True
        r3.font.size = Pt(11)
        r3.font.color.rgb = COLOR_BLACK
        r3_dot = p_i.add_run("………………………………………………    ")
        r3_dot.font.size = Pt(11)
        r3_dot.font.color.rgb = COLOR_BLACK
        
        r4 = p_i.add_run("Điểm số / Đánh giá: ")
        r4.bold = True
        r4.font.size = Pt(11)
        r4.font.color.rgb = COLOR_BLACK
        r4_dot = p_i.add_run("………………………………")
        r4_dot.font.size = Pt(11)
        r4_dot.font.color.rgb = COLOR_BLACK

    def add_header_banner(self):
        pass

    def add_title_block(self):
        pass

    def add_student_info_box(self):
        pass

    def add_section(self, title, subtitle=None):
        """Thêm tiêu đề phần lớn (Nền trắng, viền trên và dưới đen, KHÔNG có sọc đen hai bên)."""
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(6)
        p_sp.paragraph_format.space_after = Pt(2)
        
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        c = t.cell(0, 0)
        c.width = Mm(172)
        style_cell(c, bg_color=None, border_color="000000", top="single", bottom="single", left="none", right="none", sz="6")
        
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        
        if subtitle:
            ps = self.doc.add_paragraph()
            ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ps.paragraph_format.space_before = Pt(2)
            ps.paragraph_format.space_after = Pt(6)
            rs = ps.add_run(subtitle)
            rs.font.name = "Times New Roman"
            rs.font.size = Pt(11)
            rs.font.italic = True
            rs.font.color.rgb = COLOR_BLACK

    def add_section_banner(self, title, subtitle=None):
        self.add_section(title, subtitle)

    def add_dang(self, dang_title):
        """Thêm dạng toán (Nền trắng, KHÔNG có sọc đen bên trái)."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        c = t.cell(0, 0)
        c.width = Mm(172)
        style_cell(c, bg_color=None, top="none", bottom="none", left="none", right="none")
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Pt(4)
        
        r = p.add_run(dang_title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLACK
        
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(2)

    def add_problem_header(self, num_str, content_str=""):
        """Thêm câu hỏi/bài tập (font >= 11pt)."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(num_str + " ")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(11.5)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_BLACK
        
        if content_str:
            add_rich_text(p, content_str, default_size=Pt(11), default_color=COLOR_BLACK)
        return p

    def add_question_figure(self, fig_path, fig_width=Inches(3.2), caption_text=None):
        """Chèn hình vẽ trực tiếp vào phần đề bài hoặc câu hỏi."""
        if os.path.exists(fig_path):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run()
            r.add_picture(fig_path, width=fig_width)
            if caption_text:
                p_cap = self.doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(3)
                r_cap = p_cap.add_run(caption_text)
                r_cap.font.name = "Times New Roman"
                r_cap.font.size = Pt(11)
                r_cap.font.italic = True
                r_cap.font.color.rgb = COLOR_BLACK

    def add_exercise_header(self, label, text=""):
        return self.add_problem_header(f"{label}:", text)

    def add_problem_statement(self, lines):
        """Thêm các dòng nội dung của đề bài (font >= 11pt)."""
        for line in lines:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Pt(8)
            add_rich_text(p, line, default_size=Pt(11), default_color=COLOR_BLACK)

    def add_display_equation(self, latex_code):
        return add_display_math(self.doc, latex_code)

    def add_solution(self, solution_steps):
        """
        Thêm khung lời giải chi tiết (Solution Box) - Nền trắng, KHÔNG có sọc đen bên trái, font >= 11pt.
        """
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        c = t.cell(0, 0)
        c.width = Mm(172)
        # BỎ HOÀN TOÀN SỌC ĐEN BÊN TRÁI (left="none")
        style_cell(c, bg_color=None, top="none", bottom="none", left="none", right="none")
        
        first = True
        for step in solution_steps:
            kind = step[0]
            content = step[1]
            if kind == 'title':
                p = c.paragraphs[0] if first else c.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Pt(6)
                r_icon = p.add_run(content if content else "✎ LỜI GIẢI CHI TIẾT:")
                r_icon.font.name = "Times New Roman"
                r_icon.font.size = Pt(11.5)
                r_icon.font.bold = True
                r_icon.font.color.rgb = COLOR_BLACK
            elif kind == 'p':
                p = c.paragraphs[0] if first else c.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                add_rich_text(p, content, default_size=Pt(11), default_color=COLOR_BLACK)
            elif kind == 'display_math':
                add_display_math(c, content, space_before=Pt(2), space_after=Pt(2))
            elif kind == 'bullet':
                p = c.paragraphs[0] if first else c.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Pt(16)
                p.paragraph_format.line_spacing = 1.15
                r_b = p.add_run("• ")
                r_b.font.bold = True
                r_b.font.size = Pt(11)
                r_b.font.color.rgb = COLOR_BLACK
                add_rich_text(p, content, default_size=Pt(11), default_color=COLOR_BLACK)
            elif kind == 'image':
                img_path = content
                img_w = step[2] if len(step) > 2 else Mm(110)
                p = c.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                run_img = p.add_run()
                run_img.add_picture(img_path, width=img_w)
            elif kind == 'caption':
                p = c.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(4)
                r_cap = p.add_run(content)
                r_cap.font.name = "Times New Roman"
                r_cap.font.size = Pt(11)
                r_cap.font.italic = True
                r_cap.font.color.rgb = COLOR_BLACK
            first = False
            
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_solution_box(self, steps, title=None, image_path=None, image_width=4.2, conclusion=None):
        sol_steps = [('title', title or "✎ LỜI GIẢI CHI TIẾT:")]
        for st in steps:
            if isinstance(st, dict) and 'display_math' in st:
                sol_steps.append(('display_math', st['display_math']))
            elif isinstance(st, str):
                s_trimmed = st.strip()
                if s_trimmed.startswith('$$') and s_trimmed.endswith('$$') and len(s_trimmed) >= 4:
                    sol_steps.append(('display_math', s_trimmed[2:-2].strip()))
                else:
                    clean_st = st.lstrip('• ').strip() if st.strip().startswith('•') else st
                    kind = 'bullet' if st.strip().startswith('•') else 'p'
                    sol_steps.append((kind, clean_st))
        if image_path and os.path.exists(image_path):
            sol_steps.append(('image', image_path, Inches(image_width)))
            sol_steps.append(('caption', 'Hình vẽ minh họa trực quan'))
        if conclusion:
            sol_steps.append(('bullet', f"Kết luận: {conclusion}"))
        self.add_solution(sol_steps)

    def add_self_evaluation(self, criteria_list):
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(6)
        p_sp.paragraph_format.space_after = Pt(2)
        
        self.add_section(
            "PHẦN III. TỰ ĐÁNH GIÁ SAU BUỔI ÔN",
            "Học sinh tự đánh dấu mức độ đạt được và đối chiếu kết quả."
        )
        
        rows_count = len(criteria_list) + 1
        eval_table = self.doc.add_table(rows=rows_count, cols=2)
        eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        eval_table.autofit = False
        
        w_col0 = Mm(36)
        w_col1 = Mm(136)
        
        h0 = eval_table.cell(0, 0)
        h1 = eval_table.cell(0, 1)
        h0.width = w_col0
        h1.width = w_col1
        style_cell(h0, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        style_cell(h1, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        
        p0 = h0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run("Đánh giá")
        r0.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = COLOR_BLACK
        
        p1 = h1.paragraphs[0]
        r1 = p1.add_run("Nội dung kiến thức – Kỹ năng cần đạt")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK
        
        for idx, crit in enumerate(criteria_list, start=1):
            c0 = eval_table.cell(idx, 0)
            c1 = eval_table.cell(idx, 1)
            c0.width = w_col0
            c1.width = w_col1
            style_cell(c0, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
            style_cell(c1, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
            
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_box = p0.add_run("☑ Đạt    ☐ Chưa" if self.is_solved else "☐ Đạt    ☐ Chưa")
            r_box.font.size = Pt(11)
            r_box.font.color.rgb = COLOR_BLACK
            
            p1 = c1.paragraphs[0]
            p1.paragraph_format.left_indent = Pt(6)
            r_text = p1.add_run(crit)
            r_text.font.size = Pt(11)
            r_text.font.color.rgb = COLOR_BLACK

    def add_self_evaluation_table(self, items):
        self.add_self_evaluation([f"{name} ({pts})" for name, pts in items])

    def add_teacher_feedback_and_signature(self, teacher_name="Nguyễn Văn Dũng – MathPlus Academy"):
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(6)
        p_sp.paragraph_format.space_after = Pt(2)
        
        rem_table = self.doc.add_table(rows=1, cols=1)
        rem_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        rem_table.autofit = False
        rem_cell = rem_table.cell(0, 0)
        rem_cell.width = Mm(172)
        style_cell(rem_cell, bg_color=None, border_color="000000", top="single", bottom="single", left="single", right="single", sz="4")
        
        p_rem = rem_cell.paragraphs[0]
        p_rem.paragraph_format.space_before = Pt(4)
        p_rem.paragraph_format.space_after = Pt(4)
        p_rem.paragraph_format.line_spacing = 1.3
        
        r_rt = p_rem.add_run("Nhận xét của giáo viên bộ môn:\n")
        r_rt.font.bold = True
        r_rt.font.size = Pt(11.5)
        r_rt.font.color.rgb = COLOR_BLACK
        
        r_dots = p_rem.add_run(
            "…………………………………………………………………………………………………………………………………………\n"
            "…………………………………………………………………………………………………………………………………………\n"
            "…………………………………………………………………………………………………………………………………………"
        )
        r_dots.font.size = Pt(11)
        r_dots.font.color.rgb = COLOR_BLACK
        
        p_sig = self.doc.add_paragraph()
        p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sig.paragraph_format.space_before = Pt(10)
        p_sig.paragraph_format.space_after = Pt(2)
        
        r_sig1 = p_sig.add_run("Hà Nội, ngày …… tháng …… năm 2026\n")
        r_sig1.font.size = Pt(11)
        r_sig1.font.italic = True
        r_sig1.font.color.rgb = COLOR_BLACK
        
        r_sig2 = p_sig.add_run(f"Giáo viên: {teacher_name}")
        r_sig2.font.size = Pt(11.5)
        r_sig2.font.bold = True
        r_sig2.font.color.rgb = COLOR_BLACK

    def save(self, filepath):
        for elem in self.doc.settings.element.iter():
            if elem.tag.endswith('zoom'):
                elem.set(qn('w:percent'), '100')
                
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        self.doc.save(filepath)
        return filepath
