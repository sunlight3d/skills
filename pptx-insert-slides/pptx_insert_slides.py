#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_insert_slides.py

A utility script to insert beautifully designed, 100% native editable PowerPoint (.pptx)
slides into an existing presentation at any specified slide position.

Key Features:
- 100% Native PowerPoint Shapes & TextFrames (No screenshots/images).
- Fully editable: Edit text, change colors, resize shapes directly in Microsoft PowerPoint.
- Theme Harmonization: Matches the color scheme (#004C7E, #0077B6, etc.) and style of existing slides.
- Typography Standard: Font size >= 22pt everywhere for high legibility and presentation standards.
- Concise Content: Clean cards and concise bullet points preventing text overflow.
- Flexible Layouts:
    * '2_column': 2 comparison/parallel cards + bottom summary card
    * 'grid_2x2': 4 categorical cards arranged in a 2x2 grid
    * 'stacked_rows': 3-4 horizontal cards with badges and descriptions
"""

import os
import sys
import json
import shutil
import argparse
from typing import List, Dict, Any, Optional

try:
    import pptx
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Error: python-pptx is required. Install via: pip install python-pptx")
    sys.exit(1)


# Default Technology Color Palette
THEME_COLORS = {
    "title": RGBColor(0, 76, 126),       # #004C7E Dark Navy Blue
    "category": RGBColor(0, 119, 182),   # #0077B6 Primary Blue
    "text_dark": RGBColor(15, 23, 42),   # #0F172A Near Black
    "text_body": RGBColor(30, 41, 59),   # #1E293B Body Charcoal
    "text_sub": RGBColor(100, 116, 139), # #64748B Secondary Slate
    "card_bg": RGBColor(248, 250, 252),  # #F8FAFC Off-white
    "blue_border": RGBColor(0, 119, 182),
    "green_border": RGBColor(21, 128, 61),  # #15803D
    "red_border": RGBColor(220, 38, 38),    # #DC2626
    "orange_border": RGBColor(234, 88, 12), # #EA580C
    "slate_border": RGBColor(51, 65, 85),   # #334155
    "sky_border": RGBColor(2, 132, 199),    # #0284C7
    "bottom_bg": RGBColor(240, 247, 255),   # #F0F7FF
    "bottom_border": RGBColor(147, 197, 253),# #93C5FD
}


def hex_to_rgb(hex_str: str) -> RGBColor:
    """Convert hex string (e.g. '#004C7E' or '004C7E') to RGBColor."""
    s = hex_str.lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    if len(s) == 6:
        r = int(s[0:2], 16)
        g = int(s[2:4], 16)
        b = int(s[4:6], 16)
        return RGBColor(r, g, b)
    return THEME_COLORS["category"]


def add_header(slide, category: str, title: str, font_name: str, width: float):
    """Add category & title header textbox at top of slide."""
    header_box = slide.shapes.add_textbox(Pt(50), Pt(24), Pt(width - 100), Pt(76))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.alignment = PP_ALIGN.CENTER
    p_cat.font.size = Pt(22)
    p_cat.font.bold = True
    p_cat.font.name = font_name
    p_cat.font.color.rgb = THEME_COLORS["category"]

    p_tit = tf.add_paragraph()
    p_tit.text = title.upper()
    p_tit.alignment = PP_ALIGN.CENTER
    p_tit.font.size = Pt(34)
    p_tit.font.bold = True
    p_tit.font.name = font_name
    p_tit.font.color.rgb = THEME_COLORS["title"]
    p_tit.space_before = Pt(4)


def build_slide_2_column(slide, slide_data: Dict[str, Any], font_name: str, width: float, height: float):
    """Build a 2-column comparison/parallel slide with an optional bottom summary card."""
    add_header(slide, slide_data.get("category", ""), slide_data.get("title", ""), font_name, width)

    has_bottom = "bottom_card" in slide_data and slide_data["bottom_card"]
    card_h = 380 if has_bottom else 560
    card_w = (width - 130) / 2 # 575pt on 1280pt canvas
    x_left = 50
    x_right = x_left + card_w + 30
    y_cards = 115

    # 1. Left Card
    left = slide_data.get("left_card", {})
    c_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x_left), Pt(y_cards), Pt(card_w), Pt(card_h))
    bg_l = hex_to_rgb(left.get("bg_color", "#F8FAFC"))
    br_l = hex_to_rgb(left.get("border_color", "#0077B6"))
    c_left.fill.solid(); c_left.fill.fore_color.rgb = bg_l
    c_left.line.color.rgb = br_l; c_left.line.width = Pt(2.5)

    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = Pt(24)
    tf_l.margin_top = tf_l.margin_bottom = Pt(16)

    p_lh = tf_l.paragraphs[0]
    p_lh.text = f"{left.get('badge', '')} {left.get('title', '')}".strip()
    p_lh.font.size = Pt(25)
    p_lh.font.bold = True
    p_lh.font.name = font_name
    p_lh.font.color.rgb = br_l

    for item in left.get("items", []):
        p = tf_l.add_paragraph()
        p.space_before = Pt(10)
        _render_card_item(p, item, font_name, br_l)

    # 2. Right Card
    right = slide_data.get("right_card", {})
    c_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x_right), Pt(y_cards), Pt(card_w), Pt(card_h))
    bg_r = hex_to_rgb(right.get("bg_color", "#F8FAFC"))
    br_r = hex_to_rgb(right.get("border_color", "#15803D"))
    c_right.fill.solid(); c_right.fill.fore_color.rgb = bg_r
    c_right.line.color.rgb = br_r; c_right.line.width = Pt(2.5)

    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = Pt(24)
    tf_r.margin_top = tf_r.margin_bottom = Pt(16)

    p_rh = tf_r.paragraphs[0]
    p_rh.text = f"{right.get('badge', '')} {right.get('title', '')}".strip()
    p_rh.font.size = Pt(25)
    p_rh.font.bold = True
    p_rh.font.name = font_name
    p_rh.font.color.rgb = br_r

    for item in right.get("items", []):
        p = tf_r.add_paragraph()
        p.space_before = Pt(10)
        _render_card_item(p, item, font_name, br_r)

    # 3. Bottom Card (if present)
    if has_bottom:
        bot = slide_data["bottom_card"]
        c_bot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(50), Pt(515), Pt(width - 100), Pt(175))
        bg_b = hex_to_rgb(bot.get("bg_color", "#F0F7FF"))
        br_b = hex_to_rgb(bot.get("border_color", "#93C5FD"))
        c_bot.fill.solid(); c_bot.fill.fore_color.rgb = bg_b
        c_bot.line.color.rgb = br_b; c_bot.line.width = Pt(2)

        tf_b = c_bot.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = Pt(24)
        tf_b.margin_top = tf_b.margin_bottom = Pt(14)

        p_bh = tf_b.paragraphs[0]
        p_bh.text = bot.get("title", "")
        p_bh.font.size = Pt(24)
        p_bh.font.bold = True
        p_bh.font.name = font_name
        p_bh.font.color.rgb = THEME_COLORS["title"]

        for item in bot.get("items", []):
            p = tf_b.add_paragraph()
            p.space_before = Pt(6)
            if isinstance(item, str):
                r = p.add_run(); r.text = item; r.font.size = Pt(22); r.font.name = font_name; r.font.color.rgb = THEME_COLORS["text_body"]
            else:
                _render_card_item(p, item, font_name, br_b)


def build_slide_grid_2x2(slide, slide_data: Dict[str, Any], font_name: str, width: float, height: float):
    """Build a 2x2 grid slide containing 4 distinct categorized cards."""
    add_header(slide, slide_data.get("category", ""), slide_data.get("title", ""), font_name, width)

    cards = slide_data.get("cards", [])
    if len(cards) < 4:
        while len(cards) < 4:
            cards.append({"title": f"Mục {len(cards)+1}", "items": []})

    card_w = (width - 130) / 2 # 575pt
    row1_y = 112
    row1_h = 280
    row2_y = 405
    row2_h = 288

    positions = [
        (50, row1_y, card_w, row1_h),
        (50 + card_w + 30, row1_y, card_w, row1_h),
        (50, row2_y, card_w, row2_h),
        (50 + card_w + 30, row2_y, card_w, row2_h)
    ]

    default_borders = ["#DC2626", "#EA580C", "#0284C7", "#334155"]
    default_bgs = ["#FEF2F2", "#FFF7ED", "#F0F9FF", "#F8FAFC"]

    for i in range(4):
        c_info = cards[i]
        x, y, w, h = positions[i]

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        bg_col = hex_to_rgb(c_info.get("bg_color", default_bgs[i]))
        br_col = hex_to_rgb(c_info.get("border_color", default_borders[i]))
        shape.fill.solid(); shape.fill.fore_color.rgb = bg_col
        shape.line.color.rgb = br_col; shape.line.width = Pt(2.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(22)
        tf.margin_top = tf.margin_bottom = Pt(14)

        # Card Title
        badge_text = c_info.get("badge", "").strip()
        card_title = c_info.get("title", "").strip()
        full_title = f"{badge_text}: {card_title}" if badge_text else card_title

        p_h = tf.paragraphs[0]
        p_h.text = full_title
        p_h.font.size = Pt(24)
        p_h.font.bold = True
        p_h.font.name = font_name
        p_h.font.color.rgb = br_col

        for item in c_info.get("items", []):
            p = tf.add_paragraph()
            p.space_before = Pt(8)
            _render_card_item(p, item, font_name, br_col)


def build_slide_stacked_rows(slide, slide_data: Dict[str, Any], font_name: str, width: float, height: float):
    """Build a stacked horizontal rows layout (3 or 4 rows)."""
    add_header(slide, slide_data.get("category", ""), slide_data.get("title", ""), font_name, width)

    rows = slide_data.get("rows", [])
    count = min(len(rows), 4)
    if count == 0:
        return

    row_h = (580 - (count - 1) * 16) / count
    y_start = 115

    default_borders = ["#0077B6", "#15803D", "#EA580C", "#DC2626"]
    default_bgs = ["#F0F7FF", "#F0FDF4", "#FFF7ED", "#FEF2F2"]

    for i in range(count):
        r_info = rows[i]
        y = y_start + i * (row_h + 16)

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(50), Pt(y), Pt(width - 100), Pt(row_h))
        bg_col = hex_to_rgb(r_info.get("bg_color", default_bgs[i % 4]))
        br_col = hex_to_rgb(r_info.get("border_color", default_borders[i % 4]))
        shape.fill.solid(); shape.fill.fore_color.rgb = bg_col
        shape.line.color.rgb = br_col; shape.line.width = Pt(2)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(24)
        tf.margin_top = tf.margin_bottom = Pt(12)

        p = tf.paragraphs[0]
        badge = r_info.get("badge", "").strip()
        title = r_info.get("title", "").strip()
        if badge:
            r1 = p.add_run(); r1.text = f"[{badge}] "; r1.font.bold = True; r1.font.size = Pt(24); r1.font.name = font_name; r1.font.color.rgb = br_col
        if title:
            r2 = p.add_run(); r2.text = f"{title}: "; r2.font.bold = True; r2.font.size = Pt(24); r2.font.name = font_name; r2.font.color.rgb = THEME_COLORS["title"]

        desc = r_info.get("desc", "") or r_info.get("text", "")
        if desc:
            r3 = p.add_run(); r3.text = desc; r3.font.size = Pt(22); r3.font.name = font_name; r3.font.color.rgb = THEME_COLORS["text_body"]


def _render_card_item(paragraph, item: Any, font_name: str, accent_color: RGBColor):
    """Render a card item with label, text, and optional highlighted runs (all font >= 22pt)."""
    if isinstance(item, str):
        r = paragraph.add_run()
        r.text = item
        r.font.size = Pt(22)
        r.font.name = font_name
        r.font.color.rgb = THEME_COLORS["text_body"]
        return

    label = item.get("label", "")
    text = item.get("text", "")
    highlight = item.get("highlight", "")

    if label:
        r_lbl = paragraph.add_run()
        r_lbl.text = label
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(22)
        r_lbl.font.name = font_name
        r_lbl.font.color.rgb = THEME_COLORS["text_dark"]

    if highlight and highlight in text:
        parts = text.split(highlight, 1)
        if parts[0]:
            r0 = paragraph.add_run(); r0.text = parts[0]; r0.font.size = Pt(22); r0.font.name = font_name; r0.font.color.rgb = THEME_COLORS["text_body"]
        rh = paragraph.add_run()
        rh.text = highlight
        rh.font.bold = True
        rh.font.size = Pt(22)
        rh.font.name = font_name
        rh.font.color.rgb = hex_to_rgb(item.get("highlight_color", "#DC2626")) if "highlight_color" in item else accent_color
        if parts[1]:
            r1 = paragraph.add_run(); r1.text = parts[1]; r1.font.size = Pt(22); r1.font.name = font_name; r1.font.color.rgb = THEME_COLORS["text_body"]
    else:
        r_txt = paragraph.add_run()
        r_txt.text = text
        r_txt.font.size = Pt(22)
        r_txt.font.name = font_name
        r_txt.font.color.rgb = THEME_COLORS["text_body"]


def insert_slides(pptx_path: str, after_slide: int, slides_data: List[Dict[str, Any]],
                  output_path: Optional[str] = None, font_name: str = "Segoe UI") -> str:
    """
    Insert native editable slides after the given 1-based slide index.
    """
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

    target_output = output_path or pptx_path

    # Auto-backup if in-place update
    if target_output == pptx_path:
        bak_path = pptx_path + ".bak"
        if not os.path.exists(bak_path):
            shutil.copy2(pptx_path, bak_path)
            print(f"Created backup: {bak_path}")

    prs = Presentation(pptx_path)
    total_slides_before = len(prs.slides)
    slide_w_pt = prs.slide_width / Pt(1)
    slide_h_pt = prs.slide_height / Pt(1)

    # Validate insertion position
    # after_slide = 10 means inserted slides start at index 10 (1-based Slide 11)
    insert_idx = max(0, min(after_slide, total_slides_before))

    blank_layout = prs.slide_layouts[6]
    sldIdLst = prs.slides._sldIdLst

    new_slide_elements = []

    for s_idx, s_data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        layout_type = s_data.get("layout", "2_column").lower()

        if layout_type in ["2_column", "2_col", "columns"]:
            build_slide_2_column(slide, s_data, font_name, slide_w_pt, slide_h_pt)
        elif layout_type in ["grid_2x2", "2x2", "grid"]:
            build_slide_grid_2x2(slide, s_data, font_name, slide_w_pt, slide_h_pt)
        elif layout_type in ["stacked_rows", "rows", "stacked"]:
            build_slide_stacked_rows(slide, s_data, font_name, slide_w_pt, slide_h_pt)
        else:
            build_slide_2_column(slide, s_data, font_name, slide_w_pt, slide_h_pt)

        # Track the newly created slide XML element
        elem = sldIdLst[-1]
        new_slide_elements.append(elem)

    # Reorder slides into position
    for elem in new_slide_elements:
        sldIdLst.remove(elem)

    for i, elem in enumerate(new_slide_elements):
        sldIdLst.insert(insert_idx + i, elem)

    prs.save(target_output)
    print(f"SUCCESS: Inserted {len(slides_data)} editable slide(s) after slide {after_slide}.")
    print(f"Output saved to: {target_output}")
    return target_output


def main():
    parser = argparse.ArgumentParser(
        description="Insert 100% native editable PowerPoint slides with matching theme, font >= 22pt, and concise text."
    )
    parser.add_argument("pptx", help="Path to input PowerPoint (.pptx) file")
    parser.add_argument("--after-slide", "-a", type=int, required=True,
                        help="1-based slide index after which new slides are inserted (e.g. 10)")
    parser.add_argument("--content", "-c", type=str,
                        help="Path to JSON file containing slides definition")
    parser.add_argument("--json-str", "-j", type=str,
                        help="Direct JSON string containing slides definition")
    parser.add_argument("--output", "-o", type=str, default=None,
                        help="Destination .pptx path (default: overwrite input in-place with .bak backup)")
    parser.add_argument("--font", "-f", type=str, default="Segoe UI",
                        help="Font family for slides (default: 'Segoe UI')")

    args = parser.parse_args()

    # Load content
    slides_data = None
    if args.content:
        with open(args.content, "r", encoding="utf-8") as f:
            slides_data = json.load(f)
    elif args.json_str:
        slides_data = json.loads(args.json_str)
    else:
        # Check standard input
        if not sys.stdin.isatty():
            slides_data = json.load(sys.stdin)
        else:
            print("Error: Either --content <file.json> or --json-str '<json>' must be provided.")
            sys.exit(1)

    if isinstance(slides_data, dict):
        slides_data = [slides_data]

    insert_slides(
        pptx_path=args.pptx,
        after_slide=args.after_slide,
        slides_data=slides_data,
        output_path=args.output,
        font_name=args.font
    )


if __name__ == "__main__":
    main()
