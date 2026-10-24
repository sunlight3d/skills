import os
import sys
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
import pymupdf
from PIL import Image
import io

def convert_pdf_to_docx(pdf_path, docx_output_path=None, max_pages=None, dpi=150):
    if docx_output_path is None:
        base, _ = os.path.splitext(pdf_path)
        docx_output_path = base + ".docx"

    print(f"[*] Processing PDF: {pdf_path}")
    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    print(f"[*] Total pages: {total_pages}")

    num_pages = total_pages if max_pages is None else min(total_pages, max_pages)

    # Check if PDF has native selectable text
    has_native_text = False
    sample_text = ""
    for i in range(min(5, num_pages)):
        sample_text += doc[i].get_text().strip()
    if len(sample_text) > 200:
        has_native_text = True

    out_doc = Document()
    
    # Configure document style
    style = out_doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    if has_native_text:
        print("[*] PDF contains native text. Extracting text directly...")
        for page_idx in range(num_pages):
            page = doc[page_idx]
            page_text = page.get_text()
            if page_text.strip():
                p_header = out_doc.add_paragraph()
                r_h = p_header.add_run(f"--- TRANG {page_idx + 1} ---")
                r_h.bold = True
                r_h.font.color.rgb = RGBColor(0x19, 0x79, 0x35)
                
                for line in page_text.split('\n'):
                    line_str = line.strip()
                    if line_str:
                        out_doc.add_paragraph(line_str)
    else:
        print("[*] PDF is scanned image-based. Running RapidOCR on pages...")
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()

        for page_idx in range(num_pages):
            page = doc[page_idx]
            pix = page.get_pixmap(dpi=dpi)
            img_bytes = pix.tobytes("png")
            
            result, elapse = engine(img_bytes)
            
            p_header = out_doc.add_paragraph()
            r_h = p_header.add_run(f"--- TRANG {page_idx + 1} ---")
            r_h.bold = True
            r_h.font.color.rgb = RGBColor(0x19, 0x79, 0x35)

            if result:
                for line_box, text, score in result:
                    if text.strip():
                        out_doc.add_paragraph(text.strip())
            else:
                out_doc.add_paragraph("[Không phát hiện văn bản trên trang này]")
            
            if (page_idx + 1) % 10 == 0 or (page_idx + 1) == num_pages:
                print(f"[*] OCR completed {page_idx + 1}/{num_pages} pages...")

    out_doc.save(docx_output_path)
    print(f"[OK] Converted PDF to DOCX successfully: {docx_output_path}")
    return docx_output_path

if __name__ == "__main__":
    if len(sys.argv) > 1:
        pdf_file = sys.argv[1]
        out_file = sys.argv[2] if len(sys.argv) > 2 else None
        convert_pdf_to_docx(pdf_file, out_file)
    else:
        print("Usage: python pdf_converter.py <input.pdf> [output.docx]")
