import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=90, bottom=90, left=110, right=110):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="A3C9A8", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_compliance_report(output_paths=None, items_data=None):
    if output_paths is None:
        output_paths = ["reviews3.docx"]

    doc = Document()

    # Chuẩn khổ giấy A4 ngang (Landscape)
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)   # 297 mm
        section.page_height = Inches(8.27)   # 210 mm
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles - Toàn bộ chữ màu đen (#000000)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Title - Chữ màu đen (#000000)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(14)
    run_title = p_title.add_run("BÁO CÁO RÀ SOÁT TUÂN THỦ QUY TRÌNH QUẢN TRỊ DỮ LIỆU")
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Section I: Tổng quan kết quả
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(4)
    r1 = h1.add_run("I. TỔNG QUAN KẾT QUẢ ĐÁNH GIÁ")
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.2
    p_intro.add_run(
        "Báo cáo tiến hành rà soát, đối chiếu toàn diện khung quy trình quản trị dữ liệu dựa trên hệ thống các văn bản quy định và khung kiến trúc hiện hành:\n"
        "- Quyết định số 1323/QĐ-BTC ban hành Khung kiến trúc và Khung quản trị dữ liệu Bộ Tài chính (1323.pdf);\n"
        "- Quyết định số 2839/QĐ-BTC ban hành Kiến trúc Cơ sở dữ liệu quốc gia về Tài chính (2839_QD_BTC kien truc CSDLQG TC.docx);\n"
        "- Quyết định số 292/QĐ-BKHCN ban hành Khung kiến trúc Chính phủ số Việt Nam, phiên bản 4.0 (292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx)."
    )

    p_stats = doc.add_paragraph()
    p_stats.paragraph_format.space_before = Pt(2)
    p_stats.paragraph_format.space_after = Pt(8)
    p_stats.paragraph_format.line_spacing = 1.2
    
    r = p_stats.add_run("• Tổng số nội dung rà soát: ")
    r.bold = True
    p_stats.add_run("15 nội dung (11 quy trình nghiệp vụ Mức 3 và 4 yêu cầu về cấu trúc khung quy trình).\n")

    r = p_stats.add_run("• Số mục Đạt yêu cầu: ")
    r.bold = True
    r_pass = p_stats.add_run("11 mục ")
    r_pass.bold = True
    r_pass.font.color.rgb = RGBColor(0x19, 0x79, 0x35) # Xanh #197935
    p_stats.add_run("(Các quy trình thành phần đã được định hình phù hợp với các văn bản quy định).\n")

    r = p_stats.add_run("• Số mục Chưa đạt / Cần hoàn thiện: ")
    r.bold = True
    r_fail = p_stats.add_run("4 mục ")
    r_fail.bold = True
    r_fail.font.color.rgb = RGBColor(0xC6, 0x28, 0x28) # Đỏ
    p_stats.add_run("(Còn thiếu nhóm quản lý dữ liệu chủ, quy trình thay đổi kiến trúc, phân công vai trò Data Owner/Steward và tài liệu quy trình chi tiết).")

    # Section II: Bảng chi tiết
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("II. BẢNG CHI TIẾT ĐỐI CHIẾU VÀ ĐÁNH GIÁ TUÂN THỦ")
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    if items_data is None:
        items_data = [
            (
                "1",
                "1.1. Quy trình Quản lý siêu dữ liệu và từ điển dữ liệu",
                "Đạt",
                [
                    (
                        "Quy trình xây dựng, cập nhật, công bố từ điển dữ liệu; trách nhiệm của đơn vị quản lý siêu dữ liệu; chuẩn hóa thuật ngữ.",
                        "1323.pdf"
                    ),
                    (
                        "b) Xây dựng Khung quản trị, quản lý dữ liệu chi tiết áp dụng đối với cơ sở dữ liệu do đơn vị mình quản lý. Thực hiện các chức năng quản trị, quản lý dữ liệu tại đơn vị, bao gồm làm sạch, quản lý chất lượng dữ liệu nguồn, cung cấp dữ liệu nguồn thông qua các công cụ và quy chế tích hợp dữ liệu, tham gia trong quá trình triển khai các mảng dữ liệu mới có liên quan đến nghiệp vụ của đơn vị mình, phối hợp với Cục Công nghệ thông tin và chuyển đổi số xây dựng từ điển dữ liệu, tham gia quản lý trong mô hình dữ liệu và khai thác dữ liệu có liên quan;",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    )
                ]
            ),
            (
                "2",
                "2.1. Quy trình Quản lý vòng đời dữ liệu",
                "Đạt",
                [
                    (
                        "a) Các chính sách liên quan đến việc quản trị vòng đời của dữ liệu, trong đó quy định rõ vai trò, trách nhiệm của các bên liên quan đến việc cung cấp, xử lý, làm sạch, tích hợp, thu nhận dữ liệu, lưu trữ, phân tích và phân phối thông tin, dữ liệu. Xây dựng và ban hành các quy định của Bộ Tài chính về: Bộ các chỉ số phục vụ công tác chỉ đạo, điều hành của Bộ Tài chính; Danh mục cơ sở dữ liệu chuyên ngành của Bộ Tài chính; Danh mục dữ liệu mở của Bộ Tài chính; Chiến lược dữ liệu ngành Tài chính đến năm 2030;",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Thời gian lưu trữ cho từng loại dữ liệu; chính sách di chuyển giữa các vùng lưu trữ (hot/warm/cold); quy trình hủy dữ liệu.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "3",
                "2.2. Quy trình Quản lý lưu trữ và hủy dữ liệu",
                "Đạt",
                [
                    (
                        "Thời gian lưu trữ cho từng loại dữ liệu; chính sách di chuyển giữa các vùng lưu trữ (hot/warm/cold); quy trình hủy dữ liệu.",
                        "1323.pdf"
                    ),
                    (
                        "Rà soát các cơ chế chính sách làm cơ sở xây dựng quy chế, quy định quản trị dữ liệu đảm bảo việc thu nhận, xử lý làm sạch, lưu trữ, phân tích và phân phối các dữ liệu phi cấu trúc và bán cấu trúc trong lĩnh vực tài chính",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    )
                ]
            ),
            (
                "4",
                "2.3. Quy trình Quản lý kho dữ liệu và báo cáo",
                "Đạt",
                [
                    (
                        "e) Kho dữ liệu dùng chung và Nền tảng tổng hợp, phân tích dữ liệu (hoặc Nền tảng phân tích, xử lý dữ liệu tổng hợp) cấp bộ, cấp tỉnh",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    ),
                    (
                        "- Cấu phần xây dựng và triển khai cơ sở dữ liệu tổng hợp về Tài chính theo kiến trúc Data Hub. Hình thành các kho dữ liệu chủ, dữ liệu nghiệp vụ trong Tổng kho dữ liệu và các kho dữ liệu dùng chung, kho dữ liệu mở, kho dữ liệu gia tăng; Thực hiện trang bị công cụ phân tích dữ liệu bao gồm công cụ báo cáo thông minh (BI) và sử dụng công nghệ trí tuệ nhân tạo (AI);",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Giám sát việc quản lý kho dữ liệu và phân tích dữ liệu",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "5",
                "3.1. Quy trình Phát hiện và xử lý lỗi chất lượng dữ liệu",
                "Đạt",
                [
                    (
                        "- Cấu phần Quản trị dữ liệu (Data Governance) là cấu phần quan trọng trong việc đảm bảo chất lượng dữ liệu và cung cấp một nguồn dữ liệu nhất quán. Dữ liệu được thu thập và tích hợp vào Cơ sở dữ liệu quốc gia về Tài chính phải được quản lý đúng cách với dữ liệu đặc tả và quản lý chất lượng dữ liệu. Dữ liệu được chuẩn hóa và đảm bảo tuân thủ các quy định về quản lý nhà nước về dữ liệu;",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Xác định, phân tích nguyên nhân, khắc phục và ngăn ngừa tái diễn các lỗi dữ liệu.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "6",
                "4.1. Quy trình Cấp quyền truy cập dữ liệu",
                "Đạt",
                [
                    (
                        "Hỗ trợ quản lí công tác định danh và xác thực điện tử các đối tượng kết nối với các hệ thống thông tin của cơ quan",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    ),
                    (
                        "Đảm bảo việc cấp, thay đổi, thu hồi quyền truy cập dữ liệu được thực hiện đúng thẩm quyền, an toàn.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "7",
                "4.2. Quy trình Ứng cứu sự cố dữ liệu",
                "Đạt",
                [
                    (
                        "Hỗ trợ công tác phản ứng, khắc phục các sự cố mất an toàn thông tin mạng của các hệ thống thông tin của cơ quan",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    ),
                    (
                        "Phát hiện, cô lập, khắc phục, khôi phục dữ liệu khi có sự cố (mất mát, rò rỉ, hỏng hóc).",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "8",
                "4.3. Quy trình Phân loại và gắn nhãn dữ liệu",
                "Đạt",
                [
                    (
                        "c) Tăng cường khả năng chuẩn hóa, bảo đảm an toàn thông tin mạng, an ninh mạng, bảo vệ dữ liệu cá nhân trong triển khai Chính phủ số.",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    ),
                    (
                        "Hướng dẫn phân loại dữ liệu theo mức độ nhạy cảm (bí mật nhà nước, dữ liệu cá nhân, dữ liệu nội bộ…); chính sách mã hóa, che giấu, kiểm soát truy cập.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "9",
                "5.1. Quy trình Công bố dữ liệu mở",
                "Đạt",
                [
                    (
                        "Cổng dữ liệu quốc gia: Là điểm đầu mối truy cập thông tin, dữ liệu có địa chỉ data.gov.vn trên mạng Internet phục vụ việc công bố dữ liệu mở, cung cấp thông tin về chia sẻ dữ liệu của cơ quan nhà nước; cung cấp các tài liệu, dịch vụ, công cụ, ứng dụng xử lý, khai thác dữ liệu được công bố bởi các cơ quan nhà nước.",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    ),
                    (
                        "Cổng dữ liệu Bộ tài chính là điểm đầu mối truy cập thông tin, dữ liệu trên môi trường mạng phục vụ việc công bố dữ liệu mở, cung cấp thông tin về chia sẻ dữ liệu của cơ quan nhà nước trong nội bộ của Bộ Tài chính; cung cấp các tài liệu, dịch vụ, công cụ, ứng dụng xử lý, khai thác dữ liệu được công bố bởi các cơ quan nhà nước trong Bộ Tài chính.",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Đảm bảo việc công bố dữ liệu mở được thực hiện đúng quy định, kịp thời, có kiểm soát.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "10",
                "6.1. Quy trình Đăng ký và phê duyệt kết nối chia sẻ dữ liệu",
                "Đạt",
                [
                    (
                        "đ) Chủ động, phối hợp với Cục Công nghệ thông tin và chuyển đổi số xây dựng, triển khai, phát triển, nâng cấp các ứng dụng, hệ thống thông tin, cơ sở dữ liệu chuyên ngành thuộc phạm vi quản lý của đơn vị; liên thông, tích hợp, đồng bộ, chia sẻ dữ liệu nguồn với Cơ sở dữ liệu quốc gia về Tài chính;",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Kiểm soát các yêu cầu kết nối, đảm bảo tuân thủ tiêu chuẩn kỹ thuật và bảo mật.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "11",
                "6.2. Quy trình Vận hành và Quản lý sự cố tích hợp",
                "Đạt",
                [
                    (
                        "c) Quản lý vận hành hệ thống theo mô hình DataOps (DataOps là phương pháp tổ chức quản lý dữ liệu, xây dựng và áp dụng nguyên tắc để tự động hóa, tích hợp và tối ưu hóa quy trình xử lý dữ liệu, nhằm tăng tốc độ, chất lượng và hiệu quả, giảm thiểu gián đoạn hệ thống trong việc khai thác và cung cấp thông tin cho công tác quản lý, chỉ đạo điều hành);",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Xác định quy trình tích hợp dữ liệu theo các luồng tích hợp dữ liệu có kiểm soát đầu vào (data ingestion), bao gồm xác thực nguồn, kiểm tra định dạng, kiểm tra dữ liệu trùng và gắn siêu dữ liệu (metadata) nhận dạng.",
                        "1323.pdf"
                    ),
                    (
                        "6. Nền tảng tích hợp, chia sẻ dữ liệu quốc gia",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    )
                ]
            ),
            (
                "12",
                "Thiếu nhóm quy trình Quản lý dữ liệu chủ (Master Data Management)",
                "Chưa đạt",
                [
                    (
                        "- Cấu phần xây dựng và triển khai cơ sở dữ liệu tổng hợp về Tài chính theo kiến trúc Data Hub. Hình thành các kho dữ liệu chủ, dữ liệu nghiệp vụ trong Tổng kho dữ liệu và các kho dữ liệu dùng chung, kho dữ liệu mở, kho dữ liệu gia tăng; Thực hiện trang bị công cụ phân tích dữ liệu bao gồm công cụ báo cáo thông minh (BI) và sử dụng công nghệ trí tuệ nhân tạo (AI);",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Xác định các thực thể dữ liệu chủ; phân công đơn vị chủ quản; quy trình tạo mới, cập nhật, hủy; trách nhiệm của Chủ sở hữu dữ liệu (Data Owner) và Cán bộ quản trị dữ liệu (Data Steward).",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "13",
                "Thiếu Quy trình Quản lý thay đổi kiến trúc dữ liệu",
                "Chưa đạt",
                [
                    (
                        "Quản lý các đề xuất thay đổi mô hình, cấu trúc dữ liệu, đảm bảo tính nhất quán và phù hợp với chiến lược.",
                        "1323.pdf"
                    ),
                    (
                        "Kiểm tra hiệu quả của các phương án bảo đảm an toàn thông tin mạng, an ninh mạng",
                        "292_QD_KHCN khung kien truc Chinh phu so_phuluc.docx"
                    )
                ]
            ),
            (
                "14",
                "Thiếu Quy định phân công trách nhiệm quản trị dữ liệu (Data Owner/Steward)",
                "Chưa đạt",
                [
                    (
                        "b) Đào tạo, nâng cao năng lực đội ngũ cán bộ làm công tác dữ liệu ngành Tài chính cho từng vị trí thực hiện chức năng về quản trị dữ liệu.",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Xác định rõ vai trò Chủ sở hữu dữ liệu (Data Owner), Cán bộ quản trị dữ liệu (Data Steward), Cán bộ phụ trách kỹ thuật về dữ liệu (Data Custodian) ở từng đơn vị; cơ chế phối hợp với đơn vị đầu mối.",
                        "1323.pdf"
                    )
                ]
            ),
            (
                "15",
                "Thiếu mô tả bước thực hiện quy trình chi tiết (mới có sơ đồ khung)",
                "Chưa đạt",
                [
                    (
                        "Rà soát các cơ chế chính sách làm cơ sở xây dựng quy chế, quy định quản trị dữ liệu đảm bảo việc thu nhận, xử lý làm sạch, lưu trữ, phân tích và phân phối các dữ liệu phi cấu trúc và bán cấu trúc trong lĩnh vực tài chính",
                        "2839_QD_BTC kien truc CSDLQG TC.docx"
                    ),
                    (
                        "Xây dựng và hoàn thiện chi tiết Quy định, quy trình và tiêu chuẩn dữ liệu :",
                        "1323.pdf"
                    )
                ]
            )
        ]

    # Total rows = header + sum of sub-rows
    total_rows = 1 + sum(len(cits) for _, _, _, cits in items_data)
    table = doc.add_table(rows=total_rows, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="A3C9A8", sz="4")

    # Column widths: Thông tin trích dẫn là rộng nhất (4.9 inches), Tên file 1.9 inches
    col_widths = [Inches(0.45), Inches(2.1), Inches(0.75), Inches(4.9), Inches(1.9)]
    headers = ["STT", "Nội dung quy trình", "Trạng thái", "Thông tin trích dẫn", "Tên file trích dẫn"]

    # Header bảng màu #197935
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].width = col_widths[idx]
        set_cell_background(hdr_cells[idx], "197935")
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=90, right=90)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Populate rows with merging for columns 0, 1, 2
    current_row_idx = 1
    for item_idx, (stt, name, status, citations) in enumerate(items_data):
        num_cits = len(citations)
        start_row_idx = current_row_idx
        is_fail = (status == "Chưa đạt")
        bg_color = "FFF5F5" if is_fail else ("F8FAF8" if item_idx % 2 == 1 else "FFFFFF")

        # First sub-row
        first_row_cells = table.rows[start_row_idx].cells
        
        # Populate STT - Chữ màu đen (#000000)
        p_stt = first_row_cells[0].paragraphs[0]
        p_stt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_stt.paragraph_format.space_before = Pt(0)
        p_stt.paragraph_format.space_after = Pt(0)
        r_stt = p_stt.add_run(stt)
        r_stt.font.size = Pt(9.5)
        r_stt.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Populate Name - Chữ màu đen (#000000)
        p_name = first_row_cells[1].paragraphs[0]
        p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_name.paragraph_format.space_before = Pt(0)
        p_name.paragraph_format.space_after = Pt(0)
        r_name = p_name.add_run(name)
        r_name.font.size = Pt(9.5)
        r_name.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        if is_fail:
            r_name.bold = True

        # Populate Status
        p_stat = first_row_cells[2].paragraphs[0]
        p_stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_stat.paragraph_format.space_before = Pt(0)
        p_stat.paragraph_format.space_after = Pt(0)
        r_stat = p_stat.add_run(status)
        r_stat.font.size = Pt(9.5)
        r_stat.bold = True
        if is_fail:
            r_stat.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        else:
            r_stat.font.color.rgb = RGBColor(0x19, 0x79, 0x35)

        # Populate citations across sub-rows
        for cit_idx, (quote_txt, file_name) in enumerate(citations):
            row_cells = table.rows[start_row_idx + cit_idx].cells
            
            # Width and background for all cells in row
            for c_i in range(5):
                row_cells[c_i].width = col_widths[c_i]
                set_cell_background(row_cells[c_i], bg_color)
                set_cell_margins(row_cells[c_i], top=70, bottom=70, left=70, right=70)

            # Col 3: Quote (100% NGUYÊN VĂN, CHỮ THƯỜNG / REGULAR, KHÔNG BOLD, KHÔNG ITALIC)
            p_q = row_cells[3].paragraphs[0]
            p_q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_q.paragraph_format.space_before = Pt(0)
            p_q.paragraph_format.space_after = Pt(0)
            p_q.paragraph_format.line_spacing = 1.15
            rq = p_q.add_run(quote_txt)
            rq.font.size = Pt(9.0)
            rq.font.italic = False
            rq.bold = False
            if is_fail:
                rq.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
            else:
                rq.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            # Col 4: Filename - Chữ màu đen (#000000)
            p_f = row_cells[4].paragraphs[0]
            p_f.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_f.paragraph_format.space_before = Pt(0)
            p_f.paragraph_format.space_after = Pt(0)
            p_f.paragraph_format.line_spacing = 1.1
            rf = p_f.add_run(file_name)
            rf.font.size = Pt(8.5)
            rf.font.italic = False
            rf.bold = False
            rf.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Merge vertical cells for cols 0, 1, 2 if num_cits > 1
        if num_cits > 1:
            for sub_i in range(1, num_cits):
                table.rows[start_row_idx].cells[0].merge(table.rows[start_row_idx + sub_i].cells[0])
                table.rows[start_row_idx].cells[1].merge(table.rows[start_row_idx + sub_i].cells[1])
                table.rows[start_row_idx].cells[2].merge(table.rows[start_row_idx + sub_i].cells[2])

        current_row_idx += num_cits

    # Section III: Kiến nghị - Chữ màu đen (#000000)
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("III. KIẾN NGHỊ VÀ ĐỀ XUẤT ĐIỀU CHỈNH QUY TRÌNH")
    r3.font.size = Pt(13)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    recs = [
        "1. Bổ sung nhóm quy trình Quản lý dữ liệu chủ (Master Data): Bổ sung vào Mức 2 và quy định quy trình chi tiết Mức 3 \"Quy trình đăng ký và thay đổi dữ liệu chủ\" nhằm bảo đảm kiểm soát vòng đời dữ liệu chủ (tạo mới, điều chỉnh, hủy bỏ) đồng bộ giữa QĐ 1323 (1323.pdf) và QĐ 2839.",
        "2. Bổ sung Quy trình quản lý thay đổi kiến trúc dữ liệu: Bổ sung quy trình tiếp nhận, đánh giá và phê duyệt các đề xuất thay đổi cấu trúc/mô hình dữ liệu để duy trì tính toàn vẹn và đồng bộ theo Khung kiến trúc Chính phủ số (QĐ 292) và Khung kiến trúc dữ liệu ngành (QĐ 1323 - 1323.pdf).",
        "3. Làm rõ cơ chế phân công trách nhiệm quản trị dữ liệu: Trong các bước thực thi quy trình, cần xác định rõ vai trò và thẩm quyền của Chủ sở hữu dữ liệu (Data Owner), Cán bộ quản trị dữ liệu (Data Steward), Cán bộ kỹ thuật dữ liệu (Data Custodian) và các đơn vị thụ hưởng theo QĐ 1323 và QĐ 2839.",
        "4. Hoàn thiện tài liệu quy trình chi tiết: Chuyển đổi từ sơ đồ khung sang tài liệu quy trình tác nghiệp chi tiết (bao gồm: sơ đồ luồng công việc, các bước thực hiện 1..N, thời hạn xử lý từng bước, đơn vị chủ trì/phối hợp và hệ thống biểu mẫu đi kèm theo tiêu chuẩn DataOps và Quản trị dữ liệu quốc gia)."
    ]

    for rec in recs:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_before = Pt(2)
        p_r.paragraph_format.space_after = Pt(4)
        p_r.paragraph_format.line_spacing = 1.15
        r_rec = p_r.add_run(rec)
        r_rec.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # Save output
    for path in output_paths:
        doc.save(path)
        print(f"Successfully generated 100% verbatim report at {path}")

if __name__ == "__main__":
    out_paths = [
        r"c:\code\cecomtech\KhoaHocDataPlus10b\Bai04\reviews3.docx",
        r"c:\code\cecomtech\KhoaHocDataPlus10b\reviews3.docx"
    ]
    build_compliance_report(out_paths)
