import os
import sys
import json
import re
import time
import shutil
import argparse
import traceback
from pathlib import Path
from copy import copy
from difflib import SequenceMatcher
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ollama import Client

# Thêm thư mục hiện tại vào sys.path để import các module tiện ích
DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, DIR)

from utils.config_loader import load_app_config
from utils.file_handler import extract_archives, read_full_code

def flatten_nested_folders(target_dir):
    """Làm phẳng các thư mục lồng nhau đơn lẻ trong thư mục bài làm."""
    if not os.path.isdir(target_dir):
        return 0
    subdirs = [d for d in os.listdir(target_dir) if os.path.isdir(os.path.join(target_dir, d)) and not d.startswith('.')]
    flattened_count = 0
    for ti_name in sorted(subdirs):
        ti_path = os.path.join(target_dir, ti_name)
        while True:
            ds_store = os.path.join(ti_path, '.DS_Store')
            if os.path.exists(ds_store):
                try: os.remove(ds_store)
                except: pass
            
            contents = [c for c in os.listdir(ti_path) if c != '.DS_Store']
            if len(contents) == 1 and os.path.isdir(os.path.join(ti_path, contents[0])):
                x_name = contents[0]
                x_path = os.path.join(ti_path, x_name)
                x_ds_store = os.path.join(x_path, '.DS_Store')
                if os.path.exists(x_ds_store):
                    try: os.remove(x_ds_store)
                    except: pass
                for item in os.listdir(x_path):
                    src = os.path.join(x_path, item)
                    dst = os.path.join(ti_path, item)
                    if os.path.exists(dst):
                        if os.path.isdir(dst): shutil.rmtree(dst)
                        else: os.remove(dst)
                    shutil.move(src, dst)
                try: os.rmdir(x_path)
                except: shutil.rmtree(x_path, ignore_errors=True)
                flattened_count += 1
            else:
                break
    return flattened_count

class HeadlessGrader:
    def __init__(self, debai_dir, bailam_dir, template_path, ollama_url, ollama_model, api_key=None, max_workers=3):
        self.debai_dir = debai_dir
        self.bailam_dir = bailam_dir
        self.template_path = template_path
        self.ollama_model = ollama_model
        self.max_workers = max_workers
        headers = {}
        if api_key and api_key.strip():
            headers['Authorization'] = f'Bearer {api_key}'
        self.client = Client(host=ollama_url, headers=headers)
        self.rubrics = {}

    def ensure_template(self):
        default_bundled = os.path.join(DIR, "ChamBai.xlsx")
        if not os.path.exists(self.template_path):
            if os.path.exists(default_bundled):
                print(f"📋 Đang sao chép file mẫu gốc ChamBai.xlsx -> {self.template_path}")
                os.makedirs(os.path.dirname(os.path.abspath(self.template_path)), exist_ok=True)
                shutil.copy2(default_bundled, self.template_path)
            else:
                raise FileNotFoundError(f"Không tìm thấy file mẫu ChamBai.xlsx tại {self.template_path} hoặc {default_bundled}")

    def load_rubrics(self):
        self.rubrics = {}
        for f in Path(self.debai_dir).iterdir():
            if f.is_file() and f.suffix.lower() in ['.xlsx', '.xls']:
                try:
                    df = pd.read_excel(f)
                    df = df.dropna(subset=[df.columns[1]])
                    col_req = df.columns[1]
                    df[col_req] = df[col_req].astype(str).str.replace('"', "'").str.replace('\n', ' ')
                    rubric_json = df.to_json(orient='records', force_ascii=False)
                    self.rubrics[f.stem] = {
                        'path': str(f),
                        'json': rubric_json
                    }
                except Exception as e:
                    print(f"⚠️ Không đọc được rubric {f.name}: {e}")
        if not self.rubrics:
            raise ValueError(f"Không tìm thấy file rubric Excel (.xlsx) nào trong {self.debai_dir}!")
        print(f"✅ Đã nạp {len(self.rubrics)} bộ Đề bài/Rubric: {list(self.rubrics.keys())}")

    def scan_submissions(self):
        students_path = Path(self.bailam_dir)
        allowed_exts = {'.py', '.cpp', '.c', '.java', '.js', '.html', '.css', '.txt', '.sql', '.dart'}
        submissions = []
        for item in sorted(students_path.iterdir()):
            if item.name.startswith('.') or item.name in ['_mapping_cache.json', '_failed_cache.json']:
                continue
            if item.is_dir() and item.name not in ['.git', '__pycache__']:
                submissions.append((item.name, str(item), False))
            elif item.is_file() and item.suffix.lower() in allowed_exts:
                submissions.append((item.name, str(item), True))
        return submissions

    def match_subject(self, student_name, student_path, is_file):
        """Khớp bài làm của học viên với đề bài phù hợp nhất."""
        if len(self.rubrics) == 1:
            return list(self.rubrics.keys())[0]
        
        name_lower = student_name.lower().replace('-', '_').replace('.', '_').replace(' ', '_')
        for subject in self.rubrics.keys():
            sub_norm = subject.lower().replace('-', '_').replace('.', '_').replace(' ', '_')
            if sub_norm in name_lower or any(part in name_lower for part in sub_norm.split('_') if len(part) >= 3):
                return subject
        # Mặc định trả về môn đầu tiên
        return list(self.rubrics.keys())[0]

    def evaluate_with_llm(self, rubric_json, code, subject, student_name):
        prompt = f"""
Chấm điểm bài thi {subject} dựa trên Rubric JSON sau:
{rubric_json}
Bài làm học viên ({student_name}):
{code}
BỘ LUẬT CHẤM ĐIỂM (TUÂN THỦ TUYỆT ĐỐI):
1. ĐÓNG VAI GIẢNG VIÊN: Khách quan, tự nhiên.
2. CHẤM CÔNG BẰNG: BẮT BUỘC ĐÁNH GIÁ TẤT CẢ CÁC YÊU CẦU TRONG RUBRIC. Nếu học viên nộp bài trống hoặc chưa làm phần nào, PHẢI cho 0 điểm và ghi rõ lời phê. KHÔNG ĐƯỢC BỎ QUA YÊU CẦU NÀO.
3. LÀM TRÒN: 0, 0.5, 1.0... Không vượt quá Điểm chuẩn.
4. NHẬN XÉT ĐA DẠNG: BẮT BUỘC dùng tiếng Việt có dấu chuẩn, tự nhiên.
   - Nếu chưa làm: ghi ngắn gọn "Chưa làm".
   - Nếu làm sai/thiếu: ghi rõ lỗi sai cụ thể (ví dụ: "Chưa validate dữ liệu rỗng", "Sai công thức bonus").
5. BẢO VỆ JSON: KHÔNG dùng dấu ngoặc kép (") bên trong nội dung Nhận xét.

Cấu trúc JSON bắt buộc: [{{"yêu_cầu": "Tóm tắt yêu cầu từ rubric", "điểm_chấm": 0, "nhận_xét": "..."}}]
"""
        messages = [
            {'role': 'system', 'content': 'Bạn là Giảng viên chấm thi chuyên nghiệp. Chỉ xuất JSON thuần túy theo yêu cầu.'},
            {'role': 'user', 'content': prompt}
        ]
        try:
            response = self.client.chat(model=self.ollama_model, messages=messages, stream=False, format='json')
            ai_text = response['message']['content'].strip()
            match = re.search(r'(\[[\s\S]*\]|\{[\s\S]*\})', ai_text)
            clean_text = match.group(1) if match else ai_text.replace("```json", "").replace("```", "").strip()
            parsed = json.loads(clean_text)
            if isinstance(parsed, dict):
                for val in parsed.values():
                    if isinstance(val, list): return val
                return [parsed]
            return parsed
        except Exception as e:
            print(f"   ❌ Lỗi gọi LLM cho bài [{student_name}]: {e}")
            return None

    def _clone_sheet(self, src_ws, tgt_ws):
        for row in src_ws.iter_rows():
            for cell in row:
                tgt_cell = tgt_ws[cell.coordinate]
                tgt_cell.value = cell.value
                if cell.has_style:
                    tgt_cell.font, tgt_cell.border, tgt_cell.fill = copy(cell.font), copy(cell.border), copy(cell.fill)
                    tgt_cell.number_format, tgt_cell.alignment = copy(cell.number_format), copy(cell.alignment)
        if src_ws.merged_cells:
            for merged_cell in src_ws.merged_cells.ranges: tgt_ws.merge_cells(str(merged_cell))
        for col in src_ws.column_dimensions: tgt_ws.column_dimensions[col].width = src_ws.column_dimensions[col].width

    def _inject_grades(self, tgt_ws, evaluation, subject):
        def normalize_str(s):
            s = re.sub(r'[^\w\s]', '', str(s))
            return re.sub(r'\s+', ' ', s).strip().lower()

        eval_items = []
        if evaluation:
            for item in evaluation:
                req_val = str(item.get('yêu_cầu', item.get('yeu_cau', item.get('requirement', ''))))
                score_val = item.get('điểm_chấm', item.get('diem_cham', item.get('score', 0)))
                comment_val = item.get('nhận_xét', item.get('nhan_xet', item.get('comment', '')))
                try: score_val = float(score_val)
                except: score_val = 0
                eval_items.append({
                    'norm_req': normalize_str(req_val),
                    'score': score_val,
                    'comment': comment_val
                })

        col_req, col_max_score, col_score, col_review, col_comment = 2, 3, 4, 5, 6 
        header_row_idx = 1
        for row in tgt_ws.iter_rows(min_row=1, max_row=8):
            found_header = False
            for cell in row:
                val = str(cell.value).strip().lower() if cell.value else ""
                if "yêu cầu" in val or "tiêu chí" in val: col_req = cell.column; found_header = True
                elif "điểm chuẩn" in val: col_max_score = cell.column
                elif "điểm chấm" in val: col_score = cell.column
                elif "nhận xét" in val: col_comment = cell.column
            if found_header:
                header_row_idx = row[0].row
                break

        total_score, details_log, sum_row_idx = 0, [], None
        red_font = Font(color="FF0000")
        feedback_list = []

        for row in tgt_ws.iter_rows(min_row=header_row_idx + 1):
            req_cell = tgt_ws.cell(row=row[0].row, column=col_req)
            raw_req_text = str(req_cell.value).strip() if req_cell.value else ""
            req_val_lower = raw_req_text.lower()
            if req_val_lower.startswith("sum:") or req_val_lower.startswith("tổng"):
                sum_row_idx = row[0].row
                continue
            if not raw_req_text: continue

            req_normalized = normalize_str(raw_req_text)
            req_words = set(req_normalized.split())
            best_match, best_ratio = None, 0
            for item in eval_items:
                ai_req_words = set(item['norm_req'].split())
                if item['norm_req'] in req_normalized or req_normalized in item['norm_req']:
                    best_match = item
                    break
                ratio = SequenceMatcher(None, req_normalized, item['norm_req']).ratio()
                overlap = len(ai_req_words.intersection(req_words)) / len(ai_req_words) if ai_req_words else 0
                max_sim = max(ratio, overlap)
                if max_sim > best_ratio and max_sim >= 0.5:
                    best_ratio, best_match = max_sim, item

            max_score_cell = tgt_ws.cell(row=row[0].row, column=col_max_score)
            try:
                max_val = max_score_cell.value
                match = re.search(r"([0-9]*\.?[0-9]+)", str(max_val)) if isinstance(max_val, str) else None
                max_score = float(match.group(1)) if match else float(max_val or 0)
            except: max_score = 0

            short_req = (raw_req_text[:35] + '..') if len(raw_req_text) > 35 else raw_req_text.ljust(37)

            if best_match:
                eval_items.remove(best_match)
                try: score = round(float(best_match['score']) * 2) / 2
                except: score = 0
                if max_score > 0 and score > max_score: score = max_score

                tgt_ws.cell(row=row[0].row, column=col_score, value=score)
                tgt_ws.cell(row=row[0].row, column=col_review, value="")
                cmt_cell = tgt_ws.cell(row=row[0].row, column=col_comment)
                if max_score > 0 and score < max_score:
                    c_text = str(best_match.get('comment', '')).strip() or "Chưa hoàn thiện đầy đủ"
                    cmt_cell.value = c_text
                    cmt_cell.font = red_font
                    feedback_list.append(c_text)
                else:
                    cmt_cell.value = ""
                total_score += score
                details_log.append(f"   + {short_req} : {score}/{max_score}")
            else:
                tgt_ws.cell(row=row[0].row, column=col_score, value=0)
                tgt_ws.cell(row=row[0].row, column=col_review, value="")
                cmt_cell = tgt_ws.cell(row=row[0].row, column=col_comment)
                if max_score > 0:
                    cmt_cell.value = "Chưa làm"
                    cmt_cell.font = red_font
                    clean_r = re.sub(r'[\r\n\t]+', ' ', raw_req_text).strip()
                    feedback_list.append(f"Chưa làm ({clean_r[:25]})")
                else:
                    cmt_cell.value = ""
                details_log.append(f"   + {short_req} : 0.0/{max_score} (Chưa làm)")

            if cmt_cell.value:
                cmt_cell.alignment = Alignment(wrap_text=True, vertical='center')
                tgt_ws.row_dimensions[row[0].row].height = None

        tgt_ws.column_dimensions['F'].width = 45

        sum_cell_address = None
        if sum_row_idx is not None:
            col_letter = get_column_letter(col_score)
            tgt_ws.cell(row=sum_row_idx, column=col_score, value=f"=SUM({col_letter}{header_row_idx + 1}:{col_letter}{sum_row_idx - 1})")
            sum_cell_address = tgt_ws.cell(row=sum_row_idx, column=col_score).coordinate

        if total_score == 0:
            summary_comment = "Chưa làm bài hoặc không có mã nguồn liên quan."
        elif not feedback_list:
            summary_comment = "Hoàn thành tốt, đầy đủ các yêu cầu."
        else:
            summary_comment = "; ".join(feedback_list[:2])
            if len(feedback_list) > 2:
                summary_comment += f" (+{len(feedback_list)-2} mục khác)"

        tgt_ws['I10'] = subject
        return total_score, details_log, sum_cell_address, summary_comment

    def _update_result_sheet(self, wb, student_name, total_score, subject, safe_sheet_name, sum_cell_address, summary_comment=None):
        result_ws = None
        for s in wb.sheetnames:
            if "result" in s.lower() or "kết quả" in s.lower():
                result_ws = wb[s]
                break
        if not result_ws: return

        start_row = 7
        lab_header_row = 10
        for r in range(1, 15):
            val = result_ws.cell(row=r, column=2).value
            if val and str(val).strip().upper() == "LAB":
                lab_header_row = r
                start_row = r + 1
                break

        if not result_ws.cell(row=lab_header_row, column=6).value:
            result_ws.cell(row=lab_header_row, column=6, value="Lời phê / Nhận xét")
            result_ws.cell(row=lab_header_row, column=6).font = Font(bold=True)

        target_row = None
        for r in range(start_row, result_ws.max_row + 10):
            val = result_ws.cell(row=r, column=2).value
            val_str = str(val).strip() if val is not None else ""
            if val_str.startswith("Tổng số:"):
                result_ws.cell(row=r, column=2).value = None
                result_ws.cell(row=r, column=2).font = Font(bold=False)
                continue
            if val_str == student_name: target_row = r

        if target_row is None:
            for r in range(start_row, result_ws.max_row + 10):
                val = result_ws.cell(row=r, column=2).value
                if not val or str(val).strip() == "":
                    target_row = r
                    break

        result_ws.cell(row=target_row, column=2, value=student_name)
        if sum_cell_address:
            result_ws.cell(row=target_row, column=3, value=f"='{safe_sheet_name}'!{sum_cell_address}")
        else:
            result_ws.cell(row=target_row, column=3, value=total_score)
        result_ws.cell(row=target_row, column=5, value=subject)
        if summary_comment:
            cell_cmt = result_ws.cell(row=target_row, column=6, value=summary_comment)
            cell_cmt.alignment = Alignment(wrap_text=True, vertical='center')
            result_ws.row_dimensions[target_row].height = None
        result_ws.column_dimensions['F'].width = 45

        student_count, last_row = 0, start_row - 1
        for r in range(start_row, result_ws.max_row + 5):
            val = result_ws.cell(row=r, column=2).value
            val_str = str(val).strip() if val is not None else ""
            if val_str and not val_str.startswith("Tổng số:"):
                student_count += 1
                last_row = max(last_row, r)
        total_cell = result_ws.cell(row=last_row + 1, column=2, value=f"Tổng số: {student_count}")
        total_cell.font = Font(bold=True)

    def grade_all(self):
        self.ensure_template()
        self.load_rubrics()
        submissions = self.scan_submissions()
        total_subs = len(submissions)
        print(f"\n🚀 Bắt đầu chấm {total_subs} bài làm bằng model [{self.ollama_model}]...")
        
        results_summary = []

        for idx, (s_name, s_path, is_file) in enumerate(submissions, start=1):
            subject = self.match_subject(s_name, s_path, is_file)
            rubric_info = self.rubrics.get(subject)
            if not rubric_info:
                print(f"❌ [{idx}/{total_subs}] Bỏ qua [{s_name}]: Không tìm thấy rubric cho môn {subject}")
                continue

            print(f"\n[{idx}/{total_subs}] ⏳ Đang chấm bài: {s_name} (Môn: {subject})...")
            code = read_full_code(s_path, is_file)
            if not code.strip():
                code = "KHÔNG TÌM THẤY MÃ NGUỒN HOẶC BÀI LÀM TRỐNG."
                print(f"   ⚠️ Bài làm trống hoặc không có file mã nguồn hợp lệ.")

            evaluation = self.evaluate_with_llm(rubric_info['json'], code, subject, s_name)
            
            # Ghi vào file Excel ChamBai.xlsx
            wb = openpyxl.load_workbook(self.template_path)
            safe_sheet = re.sub(r'[\\/*?:\[\]]', '', s_name)[:31]
            if safe_sheet in wb.sheetnames: del wb[safe_sheet]
            tgt_ws = wb.create_sheet(title=safe_sheet)

            rubric_wb = openpyxl.load_workbook(rubric_info['path'], data_only=True)
            self._clone_sheet(rubric_wb.active, tgt_ws)
            rubric_wb.close()

            total_score, details_log, sum_addr, summary_cmt = self._inject_grades(tgt_ws, evaluation, subject)
            self._update_result_sheet(wb, s_name, total_score, subject, safe_sheet, sum_addr, summary_cmt)
            wb.save(self.template_path)
            wb.close()

            print(f"   ✅ Hoàn tất [{s_name}]: {total_score} điểm.")
            print(f"   💬 Lời phê: {summary_cmt}")
            results_summary.append({
                'name': s_name,
                'subject': subject,
                'score': total_score,
                'comment': summary_cmt
            })

        print("\n" + "="*70)
        print("🏆 BẢNG TỔNG KẾT KẾT QUẢ CHẤM BÀI")
        print("="*70)
        print(f"{'Học viên':<35} | {'Môn':<15} | {'Điểm':<6} | {'Lời phê'}")
        print("-" * 70)
        for res in results_summary:
            print(f"{res['name']:<35} | {res['subject']:<15} | {res['score']:<6} | {res['comment']}")
        print("="*70)
        print(f"💾 File kết quả đã được cập nhật tại: {self.template_path}")

def main():
    parser = argparse.ArgumentParser(description="AI AutoGrader Headless CLI - Tự động chấm bài và cập nhật ChamBai.xlsx")
    parser.add_argument("--debai", default=os.path.expanduser("~/Downloads/debai"), help="Thư mục chứa đề thi/rubric Excel")
    parser.add_argument("--bailam", default=os.path.expanduser("~/Downloads/bailam"), help="Thư mục chứa bài làm học viên")
    parser.add_argument("--template", default=os.path.expanduser("~/Downloads/ChamBai.xlsx"), help="Đường dẫn file ChamBai.xlsx")
    parser.add_argument("--model", default=None, help="Tên model Ollama dùng chấm bài")
    parser.add_argument("--host", default=None, help="Host Ollama API (ví dụ: http://127.0.0.1:11434)")
    
    args = parser.parse_args()
    config, api_key = load_app_config()
    model = args.model or config.get("grader_model", "qwen3.5:397b-cloud")
    host = args.host or config.get("host", "http://127.0.0.1:11434")

    print("="*70)
    print("🤖 APTECH CHẤM BÀI TỰ ĐỘNG (HEADLESS CLI)")
    print(f"📁 Thư mục Đề bài : {args.debai}")
    print(f"📂 Thư mục Bài làm: {args.bailam}")
    print(f"📄 File Template  : {args.template}")
    print(f"🧠 Grader Model   : {model}")
    print("="*70)

    # Bước 1: Tự động giải nén và làm phẳng thư mục bài làm nếu cần
    print("🧹 [Bước 1] Quét & dọn dẹp thư mục bài làm...")
    try: extract_archives(args.bailam)
    except Exception as e: print(f"Lỗi giải nén: {e}")
    num_flattened = flatten_nested_folders(args.bailam)
    if num_flattened > 0:
        print(f"   -> Đã làm phẳng {num_flattened} thư mục lồng nhau đơn lẻ.")

    # Bước 2: Khởi tạo Grader và tiến hành chấm
    grader = HeadlessGrader(
        debai_dir=args.debai,
        bailam_dir=args.bailam,
        template_path=args.template,
        ollama_url=host,
        ollama_model=model,
        api_key=api_key
    )
    grader.grade_all()

if __name__ == "__main__":
    main()
