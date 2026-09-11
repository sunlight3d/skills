#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
pdf_to_xml.py - Trích xuất dữ liệu từ 1 file PDF hoặc thư mục chứa PDF và chuyển đổi sang XML, JSON, XLSX.
Hỗ trợ tự động nhận diện trường, bóc tách hóa đơn/bảng hàng hóa, tùy biến tên trường.
Mỗi lần chạy tự động sinh đồng thời cả 3 định dạng: .xml, .json, .xlsx cùng tên.
- Đối với XLSX: Nếu hóa đơn có từ 2 mặt hàng trở lên, tự động tách thành các dòng riêng biệt.
- Giá trị tiền tệ: Tự động loại bỏ dấu chấm/phẩy phân cách hàng nghìn (VD: 1.800.000 -> 1800000).
"""

import os
import sys
import re
import json
import argparse
import unicodedata
from pathlib import Path
import xml.etree.ElementTree as ET
from xml.dom import minidom

# Thiết lập encoding cho stdout/stderr trên Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Kiểm tra PyMuPDF
try:
    import pymupdf  # fitz
except ImportError:
    try:
        import fitz as pymupdf
    except ImportError:
        print("[LỖI] Cần cài đặt thư viện pymupdf: pip install pymupdf", file=sys.stderr)
        sys.exit(1)

# Thử import openpyxl để xuất Excel
HAVE_OPENPYXL = False
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

# Thử import RapidOCR cho fallback quét ảnh scan
HAVE_OCR = False
ocr_engine = None
try:
    from rapidocr_onnxruntime import RapidOCR
    import numpy as np
    import cv2
    HAVE_OCR = True
except ImportError:
    HAVE_OCR = False


def get_ocr_engine():
    global ocr_engine
    if HAVE_OCR and ocr_engine is None:
        try:
            ocr_engine = RapidOCR()
        except Exception as e:
            print(f"[CẢNH BÁO] Không khởi tạo được RapidOCR: {e}", file=sys.stderr)
    return ocr_engine


def remove_accents(input_str: str) -> str:
    """Loại bỏ dấu tiếng Việt để tạo tên thẻ XML an toàn."""
    s1 = re.sub(r'[đĐ]', 'd', input_str)
    nfkd_form = unicodedata.normalize('NFKD', s1)
    return ''.join([c for c in nfkd_form if not unicodedata.combining(c)])


def slugify_xml_tag(name: str) -> str:
    """
    Chuyển đổi chuỗi bất kỳ thành tên thẻ XML hợp lệ (snake_case):
    - Bắt đầu bằng chữ cái hoặc dấu gạch dưới
    - Chỉ chứa chữ cái không dấu, số, dấu gạch dưới
    """
    clean_name = remove_accents(name).lower()
    clean_name = re.sub(r'[\(\)\[\]\{\}:;,.\/\-+\*#@!&%]+', ' ', clean_name)
    clean_name = re.sub(r'^\s*(\d+\s*)+', '', clean_name)
    clean_name = re.sub(r'^\s*muc\s+', '', clean_name)
    clean_name = re.sub(r'^\s*dieu\s+', '', clean_name)
    clean_name = re.sub(r'\s+', '_', clean_name).strip('_')
    
    if not clean_name or not (clean_name[0].isalpha() or clean_name[0] == '_'):
        clean_name = 'field_' + clean_name if clean_name else 'field'
    
    return clean_name[:40]


def clean_money_value(val: str) -> str:
    """
    Loại bỏ dấu chấm, dấu phẩy phân tách hàng nghìn trong giá trị tiền tệ.
    Ví dụ: '1.800.000' -> '1800000', '25.000.000' -> '25000000'
    """
    if not val:
        return ""
    clean = re.sub(r'[\.\s,đĐ]+', '', str(val).strip())
    return clean


def normalize_money_field(key: str, val: str) -> str:
    """Kiểm tra và chuẩn hóa số tiền nếu tên trường liên quan đến tiền tệ/chi phí."""
    if not val:
        return ""
    money_keys = ("tien", "gia", "thanh_toan", "phi", "amount", "price", "total", "cost")
    if any(k in key.lower() for k in money_keys):
        clean = clean_money_value(val)
        if clean.isdigit():
            return clean
    return val


def extract_invoice_buyer_and_items(doc, pdf_path: str):
    """
    Trích xuất chuyên sâu cho hóa đơn: Lấy nội dung từ chỗ người mua hàng đến hết danh sách hàng đã mua.
    """
    page = doc[0]
    words = page.get_text("words")
    if not words:
        return None

    page_text = page.get_text("text").lower()
    if not ("người mua" in page_text or "mua hàng" in page_text or "hóa đơn" in page_text):
        return None

    record = {}
    p = Path(pdf_path)
    record["file_name"] = p.name

    # 1. Trích xuất thông tin người mua (toạ độ y từ 240 đến 360)
    buyer_words = [w for w in words if 240 <= w[1] <= 360]
    lines = []
    curr_y = None
    curr_line = []
    for w in sorted(buyer_words, key=lambda x: (round(x[1] / 7) * 7, x[0])):
        y_band = round(w[1] / 7) * 7
        if curr_y is None or abs(y_band - curr_y) > 4:
            if curr_line:
                lines.append(" ".join([x[4] for x in curr_line]))
            curr_y = y_band
            curr_line = [w]
        else:
            curr_line.append(w)
    if curr_line:
        lines.append(" ".join([x[4] for x in curr_line]))
    
    buyer_text = "\n".join(lines)

    # Họ tên người mua hàng
    m_buyer = re.search(r'Họ tên người mua hàng:[^\S\r\n]*(.*?)(?:\r?\n|$)', buyer_text)
    record["ho_ten_nguoi_mua"] = m_buyer.group(1).strip() if m_buyer else ""

    # Tên đơn vị mua
    m_unit = re.search(r'Tên đơn vị:[^\S\r\n]*(.*?)(?:\r?\n|$)', buyer_text)
    record["ten_don_vi_mua"] = m_unit.group(1).strip() if m_unit else ""

    # Mã số thuế người mua
    m_mst = re.search(r'Mã số thuế:[^\S\r\n]*([0-9\-]+)', buyer_text)
    record["ma_so_thue_mua"] = m_mst.group(1).strip() if m_mst else ""

    # Địa chỉ người mua
    m_addr = re.search(r'Địa chỉ:[^\S\r\n]*(.*?)(?:\r?\n|$)', buyer_text)
    record["dia_chi_mua"] = m_addr.group(1).strip() if m_addr else ""

    # Mã QHNS
    m_qhns = re.search(r'Mã QHNS:[^\S\r\n]*([0-9]+)', buyer_text)
    record["ma_qhns"] = m_qhns.group(1).strip() if m_qhns else ""

    # Hình thức thanh toán
    m_pay = re.search(r'Hình thức thanh toán:[^\S\r\n]*([^\r\n]+?)(?=\s*Số tài khoản:|\r?\n|$)', buyer_text)
    record["hinh_thuc_thanh_toan"] = m_pay.group(1).strip() if m_pay else ""

    # Số tài khoản người mua nếu có
    m_acc = re.search(r'Số tài khoản:[^\S\r\n]*([0-9A-Za-z]+)', buyer_text)
    record["so_tai_khoan_mua"] = m_acc.group(1).strip() if m_acc else ""

    # 2. Trích xuất bảng danh sách hàng hóa (y từ 390 đến 535)
    stt_words = [w for w in words if 390 <= w[1] <= 535 and w[0] < 45 and w[4].isdigit() and w[4] != '6']
    stt_words.sort(key=lambda w: int(w[4]))

    items = []
    if stt_words:
        boundaries = []
        for i, sw in enumerate(stt_words):
            top = 388 if i == 0 else (stt_words[i - 1][1] + sw[1]) / 2
            bot = 535 if i == len(stt_words) - 1 else (sw[1] + stt_words[i + 1][1]) / 2
            boundaries.append((sw[4], top, bot))

        for stt, top, bot in boundaries:
            row_words = [w for w in words if top <= w[1] < bot]
            name_w = [w for w in row_words if 45 <= w[0] < 285]
            name_w.sort(key=lambda w: (round(w[1] / 5) * 5, w[0]))
            name = " ".join(w[4] for w in name_w)

            dvt_w = [w for w in row_words if 285 <= w[0] < 355]
            dvt = " ".join(w[4] for w in dvt_w)

            sl_w = [w for w in row_words if 355 <= w[0] < 410]
            sl = " ".join(w[4] for w in sl_w)

            price_w = [w for w in row_words if 410 <= w[0] < 480]
            raw_price = " ".join(w[4] for w in price_w)
            price = clean_money_value(raw_price)

            total_w = [w for w in row_words if 480 <= w[0] < 580]
            raw_total = " ".join(w[4] for w in total_w)
            total = clean_money_value(raw_total)

            items.append({
                "stt": stt,
                "ten_hang_hoa": name,
                "don_vi_tinh": dvt,
                "so_luong": sl,
                "don_gia": price,
                "thanh_tien": total
            })
    record["danh_sach_hang_hoa"] = items

    # 3. Cộng tiền hàng
    total_area_words = [w for w in words if 520 <= w[1] <= 600]
    total_text = " ".join([w[4] for w in sorted(total_area_words, key=lambda x: (round(x[1] / 7) * 7, x[0]))])
    m_total = re.search(r'Cộng tiền hàng:\s*([0-9\.,]+)', total_text)
    record["cong_tien_hang"] = clean_money_value(m_total.group(1)) if m_total else ""

    return record


def extract_text_from_pdf(pdf_path: str) -> tuple:
    """
    Trích xuất toàn bộ text từ file PDF:
    - Sử dụng PyMuPDF cho văn bản số.
    - Nếu trang không có text hoặc text quá ít (< 30 ký tự), tự động fallback dùng RapidOCR.
    """
    full_text_pages = []
    try:
        doc = pymupdf.open(pdf_path)
    except Exception as e:
        print(f"[LỖI] Không thể mở file PDF {pdf_path}: {e}", file=sys.stderr)
        return "", None

    invoice_record = extract_invoice_buyer_and_items(doc, pdf_path)

    ocr = None
    for page_idx, page in enumerate(doc):
        text = page.get_text("text").strip()
        
        if len(text) < 30 and HAVE_OCR:
            if ocr is None:
                ocr = get_ocr_engine()
            if ocr:
                try:
                    pix = page.get_pixmap(dpi=150)
                    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
                    if pix.n == 4:
                        img = cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
                    ocr_res, _ = ocr(img)
                    if ocr_res:
                        ocr_lines = [line[1] for line in ocr_res if line and len(line) > 1]
                        text = "\n".join(ocr_lines).strip()
                except Exception as ex:
                    print(f"[CẢNH BÁO] Lỗi OCR trang {page_idx+1} file {os.path.basename(pdf_path)}: {ex}", file=sys.stderr)
        
        if text:
            full_text_pages.append(text)
            
    doc.close()
    return "\n\n".join(full_text_pages), invoice_record


def parse_fields_from_text(text: str, pdf_path: str) -> dict:
    """
    Phân tích và trích xuất các trường thông tin dạng key-value từ nội dung text chung.
    """
    record = {}
    p = Path(pdf_path)
    record["file_name"] = p.name
    record["file_size_bytes"] = str(p.stat().st_size) if p.exists() else "0"

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    
    key_val_pattern = re.compile(
        r'^(?:(?:\d+[\.\)]\s*)*)?'
        r'([^\n:：]{2,50})'
        r'[:：]\s*'
        r'(.*)$'
    )

    current_tag = None
    extracted_keys = set()

    for line in lines:
        m = key_val_pattern.match(line)
        if m:
            raw_key = m.group(1).strip()
            val = m.group(2).strip()
            
            if len(raw_key.split()) <= 8 and not raw_key.lower().startswith(("chú ý", "lưu ý", "cộng hòa", "độc lập")):
                tag = slugify_xml_tag(raw_key)
                if tag and tag not in ("field", "thong_tin"):
                    val = normalize_money_field(tag, val)
                    if tag in extracted_keys:
                        if val and val not in record.get(tag, ""):
                            record[tag] = f"{record[tag]}; {val}".strip("; ")
                    else:
                        record[tag] = val
                        extracted_keys.add(tag)
                    current_tag = tag
                    continue
        
        if current_tag and record.get(current_tag) and len(line) < 200:
            if not re.match(r'^\d+[\.\)]', line) and not line.endswith(":") and len(record[current_tag]) < 500:
                record[current_tag] += " " + line
            else:
                current_tag = None
        else:
            current_tag = None

    special_extractors = [
        ("ma_so_thue", r'(?:mã số thuế|mst|ma so thue|mã số dn)[:：\s]+([0-9]{10}(?:-[0-9]{3})?)', re.I),
        ("so_cong_van", r'(?:số|so|ký hiệu)[:：\s]+([0-9A-Z\/\.\-_]+(?:\/[A-Z0-9\-_]+)+)', re.I),
        ("so_dien_thoai", r'(?:điện thoại|sđt|tel|hotline)[:：\s]+([0-9\.\s\-\+]{8,15})', re.I),
        ("email", r'([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)', 0),
        ("ngay_thang", r'(?:ngày|hà nội, ngày|tp\.hcm, ngày)\s+(\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4})', re.I),
    ]

    for tag, pattern, flags in special_extractors:
        if tag not in record or not record[tag]:
            m = re.search(pattern, text, flags)
            if m:
                record[tag] = m.group(1).strip()

    ips = re.findall(r'\b(?:\d{1,3}\.){3}\d{1,3}\b', text)
    valid_ips = []
    for ip in ips:
        parts = ip.split('.')
        if len(parts) == 4 and all(0 <= int(x) <= 255 for x in parts):
            if ip not in valid_ips:
                valid_ips.append(ip)
    if valid_ips:
        record["danh_sach_ip"] = ", ".join(valid_ips)

    if lines:
        record["tieu_de_hoac_trich_yeu"] = lines[0][:150]

    return record


def scan_pdf_inputs(input_path: str):
    input_p = Path(input_path)
    if not input_p.exists():
        raise FileNotFoundError(f"Không tìm thấy đường dẫn: {input_path}")

    pdf_files = []
    if input_p.is_file():
        if input_p.suffix.lower() == ".pdf":
            pdf_files.append(input_p)
        else:
            raise ValueError(f"Tệp không phải là định dạng PDF: {input_path}")
    elif input_p.is_dir():
        for root, _, files in os.walk(input_p):
            for file in files:
                if file.lower().endswith(".pdf") and not file.startswith("~$"):
                    pdf_files.append(Path(root) / file)

    if not pdf_files:
        raise ValueError(f"Không tìm thấy file PDF nào trong: {input_path}")

    records = []
    field_stats = {}

    print(f"[*] Tìm thấy {len(pdf_files)} file PDF. Đang quét và trích xuất dữ liệu...")
    for idx, pdf in enumerate(pdf_files, 1):
        text, invoice_record = extract_text_from_pdf(str(pdf))
        if invoice_record:
            data = invoice_record
        else:
            data = parse_fields_from_text(text, str(pdf))
        
        data["_raw_text"] = text
        records.append(data)

        for k, v in data.items():
            if k == "_raw_text":
                continue
            if k not in field_stats:
                field_stats[k] = {"count": 0, "samples": []}
            field_stats[k]["count"] += 1
            
            if isinstance(v, list):
                if v and len(field_stats[k]["samples"]) < 2:
                    first_item = v[0]
                    sample_summary = f"{len(v)} mặt hàng (VD: {first_item.get('ten_hang_hoa', '')[:30]} | {first_item.get('thanh_tien', '')})"
                    if sample_summary not in field_stats[k]["samples"]:
                        field_stats[k]["samples"].append(sample_summary)
            else:
                if v and v not in field_stats[k]["samples"] and len(field_stats[k]["samples"]) < 3:
                    field_stats[k]["samples"].append(str(v)[:60])

    sorted_fields = sorted(field_stats.items(), key=lambda item: (-item[1]["count"], item[0]))
    
    field_suggestions = []
    for tag, info in sorted_fields:
        sample_str = " | ".join(info["samples"]) if info["samples"] else "(trống)"
        field_suggestions.append({
            "tag": tag,
            "count": info["count"],
            "sample": sample_str
        })

    return records, field_suggestions


def build_xml(records: list, 
              field_mapping: dict = None, 
              root_tag: str = "records", 
              item_tag: str = "item", 
              include_raw: bool = False) -> str:
    root = ET.Element(root_tag)
    mapping = field_mapping or {}

    for idx, rec in enumerate(records, 1):
        item_elem = ET.SubElement(root, item_tag, id=str(idx), file=rec.get("file_name", f"file_{idx}"))
        
        for old_tag, value in rec.items():
            if old_tag.startswith("_"):
                continue
            if not include_raw and old_tag == "raw_content":
                continue
                
            if old_tag in mapping:
                target_tag = mapping[old_tag]
                if not target_tag:
                    continue
            else:
                target_tag = old_tag
                
            valid_tag = slugify_xml_tag(target_tag)
            if not valid_tag:
                valid_tag = "field"

            if isinstance(value, list):
                container = ET.SubElement(item_elem, valid_tag)
                item_sub_tag = mapping.get("hang_hoa", "hang_hoa")
                for sub_item in value:
                    sub_elem = ET.SubElement(container, slugify_xml_tag(item_sub_tag))
                    for sub_k, sub_v in sub_item.items():
                        mapped_sub_k = mapping.get(sub_k, sub_k)
                        if mapped_sub_k:
                            child = ET.SubElement(sub_elem, slugify_xml_tag(mapped_sub_k))
                            child.text = str(sub_v) if sub_v is not None else ""
            else:
                child = ET.SubElement(item_elem, valid_tag)
                child.text = str(value) if value is not None else ""

        if include_raw and "_raw_text" in rec:
            raw_elem = ET.SubElement(item_elem, "raw_content")
            raw_elem.text = rec["_raw_text"]

    raw_xml_bytes = ET.tostring(root, encoding="utf-8")
    dom = minidom.parseString(raw_xml_bytes)
    pretty_xml = dom.toprettyxml(indent="    ", encoding="utf-8").decode("utf-8")
    
    clean_lines = [line for line in pretty_xml.splitlines() if line.strip()]
    return "\n".join(clean_lines) + "\n"


def build_json(records: list, 
               field_mapping: dict = None, 
               include_raw: bool = False) -> str:
    """Tạo dữ liệu JSON từ danh sách bản ghi và mapping."""
    mapping = field_mapping or {}
    clean_records = []

    for idx, rec in enumerate(records, 1):
        item_dict = {"id": idx}
        for old_tag, value in rec.items():
            if old_tag.startswith("_"):
                continue
            if not include_raw and old_tag == "raw_content":
                continue

            if old_tag in mapping:
                target_tag = mapping[old_tag]
                if not target_tag:
                    continue
            else:
                target_tag = old_tag

            valid_tag = slugify_xml_tag(target_tag)

            if isinstance(value, list):
                mapped_list = []
                for sub_item in value:
                    sub_dict = {}
                    for sub_k, sub_v in sub_item.items():
                        mapped_sub_k = mapping.get(sub_k, sub_k)
                        if mapped_sub_k:
                            sub_dict[slugify_xml_tag(mapped_sub_k)] = sub_v
                    mapped_list.append(sub_dict)
                item_dict[valid_tag] = mapped_list
            else:
                item_dict[valid_tag] = value

        if include_raw and "_raw_text" in rec:
            item_dict["raw_content"] = rec["_raw_text"]

        clean_records.append(item_dict)

    return json.dumps(clean_records, ensure_ascii=False, indent=2)


def to_excel_number(val):
    """Chuyển chuỗi số nguyên thành số int để Excel không hiển thị dấu chấm."""
    if val is None or val == "":
        return ""
    val_str = str(val).strip()
    if val_str.isdigit():
        return int(val_str)
    return val_str


def build_excel(records: list, 
                field_mapping: dict = None, 
                output_xlsx_path: str = "output.xlsx") -> None:
    """
    Tạo file Excel (.xlsx) chuẩn mực:
    - Nếu 1 hóa đơn có từ 2 mặt hàng trở lên, tách thành 2 hoặc nhiều dòng tương ứng.
    - Tiền tệ không dấu chấm phân tách hàng nghìn (1.800.000 -> 1800000).
    - Sheet 1: BangKe_HoaDon (Tách dòng từng mặt hàng, gộp ô thông tin chung hóa đơn theo chiều dọc).
    - Sheet 2: DuLieu_Phang (Tách dòng và lặp lại thông tin chung, không merge ô để tiện Filter/Pivot).
    """
    if not HAVE_OPENPYXL:
        print("[CẢNH BÁO] Chưa cài openpyxl, bỏ qua bước sinh file Excel.", file=sys.stderr)
        return

    mapping = field_mapping or {}
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    headers = [
        "STT", "Tên file PDF", "Họ tên người mua", "Tên đơn vị mua", "Mã số thuế",
        "Địa chỉ", "Mã QHNS", "Hình thức thanh toán", "Số tài khoản",
        "STT Hàng", "Tên hàng hoá, dịch vụ", "Đơn vị tính", "Số lượng", "Đơn giá", "Thành tiền",
        "Cộng tiền hàng"
    ]

    # 1. Sheet 1: BangKe_HoaDon (Tách dòng mặt hàng + Merge ô thông tin hóa đơn)
    ws1 = wb.active
    ws1.title = "BangKe_HoaDon"
    ws1.append(headers)

    # 2. Sheet 2: DuLieu_Phang (Dữ liệu phẳng không merge cell, tiện Lọc/Pivot)
    ws2 = wb.create_sheet(title="DuLieu_Phang")
    ws2.append(headers)

    merge_ranges_ws1 = []

    for idx, rec in enumerate(records, 1):
        fname = rec.get("file_name", "")
        buyer_name = rec.get("ho_ten_nguoi_mua", "")
        unit_name = rec.get("ten_don_vi_mua", "")
        tax_code = rec.get("ma_so_thue_mua", "")
        address = rec.get("dia_chi_mua", "")
        qhns = rec.get("ma_qhns", "")
        payment = rec.get("hinh_thuc_thanh_toan", "")
        account = rec.get("so_tai_khoan_mua", "")
        total_amount = to_excel_number(rec.get("cong_tien_hang", ""))

        items = rec.get("danh_sach_hang_hoa", [])
        if not items:
            items = [{
                "stt": "",
                "ten_hang_hoa": "",
                "don_vi_tinh": "",
                "so_luong": "",
                "don_gia": "",
                "thanh_tien": ""
            }]

        start_r_ws1 = ws1.max_row + 1
        for item_idx, it in enumerate(items):
            item_stt = to_excel_number(it.get("stt", ""))
            item_sl = to_excel_number(it.get("so_luong", ""))
            item_price = to_excel_number(it.get("don_gia", ""))
            item_total = to_excel_number(it.get("thanh_tien", ""))

            # Row data cho Sheet 1
            row_ws1 = [
                idx if item_idx == 0 else "",
                fname if item_idx == 0 else "",
                buyer_name if item_idx == 0 else "",
                unit_name if item_idx == 0 else "",
                tax_code if item_idx == 0 else "",
                address if item_idx == 0 else "",
                qhns if item_idx == 0 else "",
                payment if item_idx == 0 else "",
                account if item_idx == 0 else "",
                item_stt,
                it.get("ten_hang_hoa", ""),
                it.get("don_vi_tinh", ""),
                item_sl,
                item_price,
                item_total,
                total_amount if item_idx == 0 else ""
            ]
            ws1.append(row_ws1)

            # Row data cho Sheet 2 (Lặp lại toàn bộ thông tin)
            row_ws2 = [
                idx, fname, buyer_name, unit_name, tax_code,
                address, qhns, payment, account,
                item_stt, it.get("ten_hang_hoa", ""),
                it.get("don_vi_tinh", ""), item_sl,
                item_price, item_total,
                total_amount
            ]
            ws2.append(row_ws2)

        end_r_ws1 = ws1.max_row
        if len(items) > 1:
            merge_ranges_ws1.append((start_r_ws1, end_r_ws1))

    # Thực hiện merge cells cho Sheet 1 đối với các hóa đơn có từ 2 mặt hàng trở lên
    # Các cột thông tin hóa đơn được gộp theo chiều dọc: 1-9 và 16
    cols_to_merge = [1, 2, 3, 4, 5, 6, 7, 8, 9, 16]
    for r_start, r_end in merge_ranges_ws1:
        for c in cols_to_merge:
            ws1.merge_cells(start_row=r_start, start_column=c, end_row=r_end, end_column=c)

    # Format styling cho cả 2 sheet
    for ws in [ws1, ws2]:
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row=row, column=col)
                c.font = data_font
                c.border = thin_border
                if col in [1, 7, 8, 10, 12, 13]:
                    c.alignment = align_center
                elif col in [14, 15, 16]:
                    c.alignment = align_right
                    # Định dạng hiển thị số nguyên không dấu phân cách
                    c.number_format = '0'
                else:
                    c.alignment = align_left

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    try:
        wb.save(output_xlsx_path)
    except PermissionError:
        fallback_xlsx = output_xlsx_path.replace(".xlsx", "_new.xlsx")
        wb.save(fallback_xlsx)
        print(f"[CẢNH BÁO] File '{output_xlsx_path}' đang được mở trong Microsoft Excel.")
        print(f"[THÀNH CÔNG] Đã lưu dữ liệu mới vào: {fallback_xlsx}")
        print("  -> Bạn hãy đóng file Excel cũ để cập nhật đè lên file chính.")
    except Exception as e:
        print(f"[LỖI] Không thể lưu file Excel: {e}", file=sys.stderr)


def interactive_rename_fields(field_suggestions: list) -> dict:
    print("\n" + "="*70)
    print(" DANH SÁCH CÁC TRƯỜNG DỮ LIỆU ĐƯỢC TỰ ĐỘNG PHÁT HIỆN")
    print("="*70)
    print(f"{'STT':<4} | {'Tên thẻ mặc định':<25} | {'Số file có':<10} | {'Giá trị mẫu'}")
    print("-" * 70)
    
    for i, f in enumerate(field_suggestions, 1):
        print(f"{i:<4} | <{f['tag']}>".ljust(30) + f" | {f['count']:<10} | {f['sample'][:40]}")
    
    print("-" * 70)
    print("[?] Bạn có muốn đổi tên thẻ XML nào không?")
    print("    - Nhấn ENTER để giữ nguyên tất cả tên thẻ mặc định.")
    print("    - Hoặc nhập số thứ tự để đổi tên (VD: '1' hoặc '1,3' hoặc 'all').")
    
    try:
        choice = input("\nLựa chọn của bạn: ").strip()
    except (EOFError, KeyboardInterrupt):
        choice = ""

    mapping = {}
    if not choice:
        print("[*] Giữ nguyên tên các trường mặc định.")
        return mapping

    indices = []
    if choice.lower() == "all":
        indices = list(range(len(field_suggestions)))
    else:
        for part in choice.split(","):
            part = part.strip()
            if part.isdigit():
                idx = int(part) - 1
                if 0 <= idx < len(field_suggestions):
                    indices.append(idx)

    for idx in indices:
        f = field_suggestions[idx]
        old_tag = f["tag"]
        try:
            new_tag = input(f"Nhập tên mới cho thẻ <{old_tag}> (hoặc gõ '-' để bỏ trường): ").strip()
        except (EOFError, KeyboardInterrupt):
            new_tag = ""
        
        if new_tag == "-":
            mapping[old_tag] = None
            print(f"  -> Đã bỏ trường: {old_tag}")
        elif new_tag:
            clean_new = slugify_xml_tag(new_tag)
            mapping[old_tag] = clean_new
            print(f"  -> Đổi <{old_tag}> thành <{clean_new}>")

    return mapping


def main():
    parser = argparse.ArgumentParser(
        description="qnet-pdf-to-xml: Chuyển đổi file hoặc thư mục PDF thành XML, JSON và XLSX."
    )
    parser.add_argument("-i", "--input", required=True, help="Đường dẫn file PDF hoặc thư mục chứa các file PDF.")
    parser.add_argument("-o", "--output", help="Đường dẫn file XML kết quả. Mặc định tự sinh trong cùng thư mục.")
    parser.add_argument("-p", "--preview", action="store_true", help="Chỉ quét và in danh sách trường đề xuất dưới dạng JSON/bảng.")
    parser.add_argument("-m", "--mapping", help="Chuỗi JSON hoặc đường dẫn file .json chứa mapping đổi tên trường: {'old_tag': 'new_tag'}.")
    parser.add_argument("-t", "--interactive", action="store_true", help="Hỏi tương tác đổi tên trường trực tiếp qua terminal.")
    parser.add_argument("--root-tag", default="records", help="Tên thẻ gốc XML (mặc định: 'records').")
    parser.add_argument("--item-tag", default="item", help="Tên thẻ cho mỗi bản ghi PDF (mặc định: 'item').")
    parser.add_argument("--include-raw", action="store_true", help="Kèm theo toàn bộ text thô trong thẻ <raw_content>.")

    args = parser.parse_args()

    input_path = os.path.abspath(args.input)
    if not os.path.exists(input_path):
        print(f"[LỖI] Đường dẫn không tồn tại: {input_path}", file=sys.stderr)
        sys.exit(1)

    try:
        records, field_suggestions = scan_pdf_inputs(input_path)
    except Exception as e:
        print(f"[LỖI] Xử lý PDF thất bại: {e}", file=sys.stderr)
        sys.exit(1)

    # Chế độ Preview
    if args.preview:
        preview_data = {
            "total_files": len(records),
            "suggested_fields": field_suggestions
        }
        print(json.dumps(preview_data, ensure_ascii=False, indent=2))
        return

    # Xác định mapping
    field_mapping = {}
    if args.mapping:
        if os.path.exists(args.mapping):
            with open(args.mapping, "r", encoding="utf-8") as f:
                field_mapping = json.load(f)
        else:
            try:
                field_mapping = json.loads(args.mapping)
            except Exception:
                try:
                    import ast
                    field_mapping = ast.literal_eval(args.mapping)
                except Exception as e:
                    print(f"[CẢNH BÁO] Không parse được JSON mapping: {e}", file=sys.stderr)

    if args.interactive:
        field_mapping = interactive_rename_fields(field_suggestions)

    # Xác định file output (Base path cho XML, JSON, XLSX)
    if args.output:
        base_output = os.path.splitext(os.path.abspath(args.output))[0]
    else:
        if os.path.isfile(input_path):
            base_output = os.path.splitext(os.path.abspath(input_path))[0]
        else:
            base_output = os.path.join(input_path, "output_pdf_records")

    xml_path = base_output + ".xml"
    json_path = base_output + ".json"
    xlsx_path = base_output + ".xlsx"

    os.makedirs(os.path.dirname(xml_path), exist_ok=True)

    print(f"[*] Đang tạo dữ liệu xuất cho {len(records)} bản ghi...")

    # 1. Xuất file XML
    xml_content = build_xml(
        records=records,
        field_mapping=field_mapping,
        root_tag=args.root_tag,
        item_tag=args.item_tag,
        include_raw=args.include_raw
    )
    with open(xml_path, "w", encoding="utf-8") as f:
        f.write(xml_content)
    print(f"[THÀNH CÔNG] Đã xuất file XML  : {xml_path} ({len(xml_content.encode('utf-8')):,} bytes)")

    # 2. Xuất file JSON
    json_content = build_json(
        records=records,
        field_mapping=field_mapping,
        include_raw=args.include_raw
    )
    with open(json_path, "w", encoding="utf-8") as f:
        f.write(json_content)
    print(f"[THÀNH CÔNG] Đã xuất file JSON : {json_path} ({len(json_content.encode('utf-8')):,} bytes)")

    # 3. Xuất file XLSX
    try:
        build_excel(records=records, field_mapping=field_mapping, output_xlsx_path=xlsx_path)
        print(f"[THÀNH CÔNG] Đã xuất file Excel: {xlsx_path}")
    except Exception as e:
        print(f"[CẢNH BÁO] Lỗi khi xuất Excel: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
