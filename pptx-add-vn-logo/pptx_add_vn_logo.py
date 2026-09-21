# /// script
# dependencies = [
#   "google-auth",
#   "google-api-python-client",
#   "python-pptx",
#   "pillow",
# ]
# ///

import socket
# Force IPv4 to bypass broken IPv6 resolution on macOS / some networks
orig_getaddrinfo = socket.getaddrinfo
def forced_getaddrinfo(*args, **kwargs):
    args = list(args)
    if len(args) > 2:
        args[2] = socket.AF_INET
    else:
        while len(args) < 3:
            args.append(None)
        args[2] = socket.AF_INET
    return orig_getaddrinfo(*args, **kwargs)
socket.getaddrinfo = forced_getaddrinfo

import os
import re
import sys
import uuid
import argparse
from pathlib import Path

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass
from google.oauth2 import service_account
from googleapiclient.discovery import build

# URL to the uploaded logo image (publicly accessible on GitHub raw)
DEFAULT_IMAGE_URL = 'https://raw.githubusercontent.com/sunlight3d/skills/master/pptx-add-vn-logo/logo_vn.jpg'
LOCAL_LOGO_NAME = 'logo_vn.jpg'

# Dimensions for Logo VN:
# Original aspect ratio: 1200 x 339
IMG_WIDTH_EMU = 2400000
IMG_HEIGHT_EMU = int(IMG_WIDTH_EMU * (339 / 1200))  # 678,000 EMU
MARGIN_RIGHT = 250000
MARGIN_TOP = 200000

def get_script_dir():
    return Path(__file__).resolve().parent

def extract_id_from_url(url_or_id):
    """Extract Google Slides or Drive ID from a URL or raw ID."""
    url_or_id = url_or_id.strip().strip("'\"")
    m_slide = re.search(r'/presentation/d/([a-zA-Z0-9_-]+)', url_or_id)
    if m_slide:
        return m_slide.group(1), 'slides'
    m_folder = re.search(r'/drive/(?:u/\d+/)?folders/([a-zA-Z0-9_-]+)', url_or_id)
    if m_folder:
        return m_folder.group(1), 'folder'
    return url_or_id, 'unknown'

def find_credentials(custom_path=None):
    if custom_path and os.path.exists(custom_path):
        return custom_path
    
    script_dir = get_script_dir()
    candidates = [
        script_dir / 'credentials.json',
        Path(r'C:\code\skills\mcp-google-sheets\credentials.json'),
        Path(r'C:\Users\nguye\.gemini\config\skills\pptx-add-imas-logo\credentials.json'),
        Path(r'C:\code\skills\pptx-add-imas-logo\credentials.json'),
        Path('/Users/hoangnd/Documents/connect-gemini-api-471309-f8bf3d1a89f5.json'),
        Path('/Users/hoangnd/Documents/funix-auto-sheet-f464a0b5957e.json'),
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    return None

# -------------------------------------------------------------
# Google Slides Online Mode
# -------------------------------------------------------------
def process_google_slides(slides_service, presentation_id, image_url, remove_only=False):
    try:
        presentation = slides_service.presentations().get(presentationId=presentation_id).execute()
    except Exception as e:
        print(f"  Error fetching presentation {presentation_id}: {e}")
        return
        
    slides = presentation.get('slides', [])
    title = presentation.get('title', 'Untitled')
    page_size = presentation.get('pageSize', {})
    slide_width = page_size.get('width', {}).get('magnitude', 16256000)
    slide_height = page_size.get('height', {}).get('magnitude', 9144000)

    pos_x = slide_width - IMG_WIDTH_EMU - MARGIN_RIGHT
    pos_y = MARGIN_TOP

    print(f"  Title: {title}")
    print(f"  Total slides: {len(slides)}")
    print(f"  Slide dimensions: {slide_width}x{slide_height} EMU")
    print(f"  Logo target pos: X={pos_x}, Y={pos_y}, W={IMG_WIDTH_EMU}, H={IMG_HEIGHT_EMU}")

    requests = []
    for slide in slides:
        slide_id = slide.get('objectId')
        
        # Delete old logos created by this script
        elements = slide.get('pageElements', [])
        for el in elements:
            el_id = el.get('objectId', '')
            if el_id.startswith('logo_vn_img_'):
                requests.append({'deleteObject': {'objectId': el_id}})
                
        if not remove_only:
            img_id = f"logo_vn_img_{slide_id}_{uuid.uuid4().hex[:8]}"
            requests.append({
                'createImage': {
                    'objectId': img_id,
                    'url': image_url,
                    'elementProperties': {
                        'pageObjectId': slide_id,
                        'size': {
                            'width': {'magnitude': IMG_WIDTH_EMU, 'unit': 'EMU'},
                            'height': {'magnitude': IMG_HEIGHT_EMU, 'unit': 'EMU'}
                        },
                        'transform': {
                            'scaleX': 1.0,
                            'scaleY': 1.0,
                            'translateX': pos_x,
                            'translateY': pos_y,
                            'unit': 'EMU'
                        }
                    }
                }
            })

    if requests:
        action_name = "Removing" if remove_only else "Adding"
        print(f"  {action_name} logo VN on {len(slides)} slides...")
        body = {'requests': requests}
        slides_service.presentations().batchUpdate(presentationId=presentation_id, body=body).execute()
        print("  Done successfully!")
    else:
        print("  No action needed.")

def get_presentations_in_folder(drive_service, folder_id):
    query = f"'{folder_id}' in parents and mimeType='application/vnd.google-apps.presentation' and trashed=false"
    results = []
    page_token = None
    while True:
        response = drive_service.files().list(q=query, spaces='drive', fields='nextPageToken, files(id, name)', pageToken=page_token).execute()
        results.extend(response.get('files', []))
        page_token = response.get('nextPageToken', None)
        if page_token is None:
            break
    return results

# -------------------------------------------------------------
# Local PPTX Offline Mode
# -------------------------------------------------------------
def process_local_pptx(file_path, logo_path, remove_only=False):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print("  [!] python-pptx not installed. Install via: pip install python-pptx")
        return False

    print(f"\n[Local PPTX] Processing: {file_path}")
    prs = Presentation(file_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    pos_x = slide_width - IMG_WIDTH_EMU - MARGIN_RIGHT
    pos_y = MARGIN_TOP

    for idx, slide in enumerate(prs.slides):
        # Remove existing logo_vn shapes if any
        shapes_to_remove = []
        for shape in slide.shapes:
            if shape.name and shape.name.startswith("logo_vn_img_"):
                shapes_to_remove.append(shape)
        
        for s in shapes_to_remove:
            sp = s._element
            sp.getparent().remove(sp)

        if not remove_only and os.path.exists(logo_path):
            pic = slide.shapes.add_picture(
                logo_path,
                Emu(pos_x),
                Emu(pos_y),
                width=Emu(IMG_WIDTH_EMU),
                height=Emu(IMG_HEIGHT_EMU)
            )
            pic.name = f"logo_vn_img_{idx}_{uuid.uuid4().hex[:6]}"

    prs.save(file_path)
    print(f"  Saved updated PPTX: {file_path}")
    return True

# -------------------------------------------------------------
# CLI Entrypoint
# -------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Add Viện Đổi mới Sáng tạo & CĐS Liên ngành (logo_vn) logo to the top right of Google Slides or PPTX.")
    parser.add_argument("target", help="Google Slides URL/ID, Google Drive Folder URL/ID, or local .pptx file/directory")
    parser.add_argument("--folder", action="store_true", help="Treat target ID as Google Drive Folder ID")
    parser.add_argument("--credentials", default=None, help="Path to Google Service Account Credentials JSON")
    parser.add_argument("--image-url", default=DEFAULT_IMAGE_URL, help="URL of the logo image to insert")
    parser.add_argument("--remove-only", action="store_true", help="Remove existing logo without adding new one")
    
    args = parser.parse_args()

    target = args.target.strip().strip("'\"")
    script_dir = get_script_dir()
    local_logo_path = str(script_dir / LOCAL_LOGO_NAME)

    # 1. Check if target is a local file or directory
    target_path = Path(target)
    if target_path.exists():
        if target_path.is_file() and target_path.suffix.lower() in ['.pptx']:
            process_local_pptx(str(target_path), local_logo_path, args.remove_only)
            return
        elif target_path.is_dir():
            pptx_files = list(target_path.glob("*.pptx"))
            print(f"Found {len(pptx_files)} PPTX files in {target_path}")
            for pf in pptx_files:
                process_local_pptx(str(pf), local_logo_path, args.remove_only)
            return

    # 2. Google Slides / Drive mode
    raw_id, detected_type = extract_id_from_url(target)
    is_folder = args.folder or (detected_type == 'folder')

    creds_path = find_credentials(args.credentials)
    if not creds_path:
        print("[-] Error: Credentials JSON file not found. Provide via --credentials <path>.")
        sys.exit(1)
        
    print(f"Using credentials: {creds_path}")
    scopes = ['https://www.googleapis.com/auth/presentations', 'https://www.googleapis.com/auth/drive.readonly']
    creds = service_account.Credentials.from_service_account_file(creds_path, scopes=scopes)
    
    slides_service = build('slides', 'v1', credentials=creds)
    
    if is_folder:
        drive_service = build('drive', 'v3', credentials=creds)
        print(f"Searching for presentations in folder {raw_id}...")
        presentations = get_presentations_in_folder(drive_service, raw_id)
        print(f"Found {len(presentations)} presentations.")
        
        for i, p in enumerate(presentations):
            print(f"\n[{i+1}/{len(presentations)}] Processing '{p['name']}' (ID: {p['id']})")
            process_google_slides(slides_service, p['id'], args.image_url, args.remove_only)
    else:
        print(f"\nProcessing Google Slides presentation (ID: {raw_id})")
        process_google_slides(slides_service, raw_id, args.image_url, args.remove_only)

if __name__ == '__main__':
    main()
